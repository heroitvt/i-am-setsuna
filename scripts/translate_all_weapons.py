import openpyxl

excel_path = "D:/Viet Hoa Game/I_am_Setsuna-parameter-Steam.xlsx"
wb = openpyxl.load_workbook(excel_path)

# =========================================================================
# 1. BraveStoryNoteMessage (Rows 71, 72)
# =========================================================================
ws_note = wb["BraveStoryNoteMessage"]
ws_note.cell(71, 3).value = "Hướng Tới Tương Lai"
ws_note.cell(72, 3).value = "Theo chân <NAME=EB_END1_01> quay ngược thời gian, <NAME=CP_0001> và <NAME=CP_0002> nhận ra họ đang ở quá khứ của <NAME=MA_0003_01>. Trận quyết chiến cuối cùng nổ ra trước <NAME=MA_0006_01>.\n\nTuy nhiên, thấu hiểu nỗi đau khổ cùng cực của kẻ thù, <NAME=CP_0002> quyết định cùng chết với <NAME=EB_END1_01>, và xin <NAME=CP_0001> hãy vung nhát kiếm kết liễu. Sau hồi dằn vặt khôn nguôi, <NAME=CP_0001> đưa ra quyết định cuối cùng của đời mình.\n\nMọi chuyện đã khép lại. Thế giới bước tiếp về phía tương lai, và những người còn sống bắt đầu hành trình của riêng mình..."

# =========================================================================
# 2. BraveStoryWeaponMessage (All 79 rows: 6 to 84)
# =========================================================================
ws_wep = wb["BraveStoryWeaponMessage"]

weapon_translations = {
    6: "Quy trình rèn thanh kiếm này được truyền đời trong bộ tộc Mặt Nạ. Mọi thành viên đều tự tay rèn thanh kiếm cho riêng mình khi bắt đầu hành nghề lính đánh thuê. Lưỡi kiếm được chế tác từ kim loại đặc biệt phản ứng với ma lực của người rèn, ngăn không cho bất kỳ ai khác sử dụng. Đổi lại, nó mang lại sự linh hoạt tuyệt vời, cân bằng hoàn hảo giữa công và thủ, lý tưởng cho chiến đấu độc hành.",
    7: "Thanh kiếm này được rèn hoàn toàn từ loại kim loại siêu cứng và siêu nhẹ mang tên Levissium. Dù Levissium rất nổi tiếng trong giới thợ rèn, nó cực kỳ khó gia công với công nghệ hiện nay và rất ít người biết đến. Người rèn thanh kiếm này tập trung tối đa vào độ tiện dụng; tuy sức tấn công khiêm tốn, nó bù đắp bằng khả năng phòng thủ xuất sắc. Trọng lượng nhẹ giúp người dùng di chuyển mau lẹ và né tránh đòn đánh bất ngờ.",
    8: "Thanh kiếm được chế tác từ đá Spritnite nguyên tố Thời Gian. Thuở xưa, một pháp sư nghiên cứu cấm thuật đã phong ấn thành công năng lượng thời gian vào vật thể, và lưỡi kiếm cấm kỵ này là kết quả của việc chuyển hóa năng lượng đó vào vũ khí. Mọi thứ bị lưỡi kiếm chém qua đều chuyển sang màu trắng băng giá trước khi tan biến vào hư không, đó cũng là nguồn gốc tên gọi của nó. Đòn đánh mang sát thương hệ Thời Gian.",
    9: "Thanh kiếm mảnh phát ra ánh hào quang hoàng kim rực rỡ. Nó được rèn từ kim loại nhiễm từ được tinh luyện nhiều lần bằng ma lực hùng mạnh nhằm gia tăng mật độ vật chất. Tác động của từ tính và ma lực tạo nên lưỡi kiếm uốn cong độc đáo. Dù mỏng manh, lưỡi kiếm cực kỳ dẻo dai và chắc khỏe, được yểm ma lực giúp không bao giờ bị mẻ. Vũ khí gây sát thương nguyên tố Quang.",
    10: "Thanh kiếm cong rực cháy trong ngọn lửa đỏ thẫm. Nó được tạo ra từ đá Spritnite chứa năng lượng hỏa có ái lực đặc biệt với kim loại. Để khai thác tối đa nguồn năng lượng này, lưỡi kiếm được tôi luyện bằng quặng Teruru tích tụ ma lực tự nhiên qua hàng trăm năm. Kiếm rất dễ sử dụng, tạo ra những nhát chém uy lực mà không tốn nhiều sức. Viên ngọc đỏ rực trên chuôi kiếm gây sát thương nguyên tố Hỏa mạnh mẽ.",
    11: "Thanh kiếm được cho là do một tổ chức từng nghiên cứu ma lực cổ đại chế tạo. Lưỡi kiếm đỏ thẫm mang sức mạnh tấn công đáng nể cùng lượng ma lực khổng lồ có thể gia tăng qua rèn đúc. Nó được tạo ra dựa trên tôn chỉ 'dùng quỷ để diệt quỷ'. Rất ít vũ khí của tổ chức này còn sót lại, độ hiếm và giá trị lịch sử khiến nó trở thành báu vật săn lùng của các nhà sưu tầm.",
    12: "Thanh kiếm này được bao bọc trong luồng hàn khí băng giá buốt lạnh. Nó được rèn bằng quặng băng vĩnh cửu khai thác từ độ sâu hàng ngàn mét dưới lòng đất. Hàn khí tỏa ra từ lưỡi kiếm làm chậm dòng máu của kẻ thù khi chém trúng, gây sát thương hệ Thủy và có khả năng làm đóng băng mục tiêu.",
    13: "Thanh kiếm được rèn theo kỹ nghệ bí truyền từ một vùng đất viễn đông cổ kính. Kiếm được tôi luyện hàng trăm lần để đạt tới độ sắc bén hoàn hảo, có thể chém đứt cả ngọn gió mà không phát ra tiếng động. Kiếm tăng mạnh tỉ lệ đánh chí mạng.",
    14: "Thanh kiếm được chế tác từ các cổ vật khai quật tại tàn tích vương triều cổ xưa. Năng lượng cổ đại còn lưu lại trên thân kiếm ban cho người sử dụng khả năng hấp thụ một phần ma lực của đối thủ sau mỗi đòn đánh trúng.",
    15: "Thanh kiếm được ngưng tụ hoàn toàn từ ma lực thuần khiết vật chất hóa. Ở chế độ Xung Lực, thanh kiếm giải phóng toàn bộ tiềm năng ma thuật, gia tăng uy lực đòn tấn công gấp bội.",
    16: "Thanh đại kiếm khổng lồ hắc ám được chế tạo từ các bộ phận của quái vật hung hãn nhất. Nguồn tà khí bên trong thanh kiếm khuấy động sự cuồng nộ của người cầm lái, đổi lấy sức tấn công vật lý vô song.",
    17: "Thanh kiếm huyền bí được tạo ra từ một chiều không gian dị giới. Lưỡi kiếm dường như tồn tại giữa các ranh giới thực tại, cho phép nó chém xuyên qua các lớp phòng ngự kiên cố nhất của kẻ thù.",
    18: "Thanh kiếm phát ra ánh sáng rực rỡ với bảy sắc cầu vồng. Tích hợp ma lực của tất cả các nguyên tố, nó có khả năng thích ứng linh hoạt và khai thác điểm yếu của mọi loại quái vật.",
    19: "Thanh kiếm cổ xưa rèn từ thời tiền sử. Thuở xa xưa, giống loài Pengy từng được tôn sùng như biểu tượng của sự may mắn, và thanh kiếm này mang theo lời chúc phúc giúp gia tăng tỉ lệ rơi vật phẩm quý.",
    20: "Cặp luân đao được đeo bởi các thiếu nữ gánh vác sứ mệnh hiến tế. Vũ khí mang theo lời nguyện cầu thanh tẩy, giúp khuếch đại ma lực trị thương và bảo vệ tâm trí người mang khỏi tà niệm.",
    21: "Luân đao được mô phỏng theo hình dáng shuriken phi tiêu huyền thoại. Thiết kế khí động học hoàn hảo giúp luân đao bay lượn với tốc độ xé gió, tấn công kẻ địch từ khoảng cách an toàn.",
    22: "Cặp luân đao với hình dáng răng cưa gồ ghề độc đáo. Khi ném ra, các cạnh răng cưa cọ xát tạo ma sát làm rách giáp đối thủ và gây sát thương liên hoàn nhiều đòn.",
    23: "Luân đao tuyệt đẹp mang hình dáng vầng trăng khuyết lấp lánh. Ánh sáng dịu dàng của ánh trăng tỏa ra từ vũ khí giúp hồi phục một lượng nhỏ MP sau mỗi đòn ném trúng đích.",
    24: "Luân đao mang hình ngôi sao năm cánh lấp lánh. Tỏa ra ánh quang thần thánh xua đuổi bóng tối, gây sát thương nguyên tố Quang mạnh mẽ lên loài quái vật hắc ám.",
    25: "Luân đao mang hình mặt đồng hồ cổ. Được chế tác bởi nghệ nhân bậc thầy, vũ khí có khả năng làm chậm dòng thời gian xung quanh mục tiêu bị trúng đòn.",
    26: "Luân đao được thắt dải ruy băng duyên dáng. Tuy vẻ ngoài có phần nữ tính và mềm mại, cấu trúc bên trong cực kỳ vững chắc và hỗ trợ tăng tốc độ nạp thanh ATB.",
    27: "Hai mũi gai nhọn như cặp sừng nhô ra từ lưỡi luân đao. Tạo ra những vết thương sâu và gây sát thương vật lý hiểm hóc khi xoay tròn với tốc độ cao.",
    28: "Luân đao được chế tác từ loại quặng đặc biệt phản ứng nhạy bén với ma lực trị thương. Càng hồi phục nhiều máu cho đồng đội, sức tấn công của vũ khí càng gia tăng.",
    29: "Luân đao phát ra ánh bạc thanh khiết. Được đúc từ bạc nguyên chất yểm bùa hộ mệnh, nó tăng cường khả năng phòng thủ và kháng cự mọi trạng thái bất lợi cho người đeo.",
    30: "Chiếc luân đao dễ thương mô phỏng hình chú chim Pengy. Nghệ nhân chế tác đã thổi vào nó niềm vui tươi, giúp người sử dụng luôn duy trì tinh thần lạc quan và tăng tỉ lệ nhận EXP.",
    31: "Vũ khí trông như một tác phẩm điêu khắc hoa tuyết hoàn mỹ. Tỏa ra hơi thở buốt giá của mùa đông vĩnh cửu, làm giảm tốc độ di chuyển của mọi kẻ thù chạm phải.",
    32: "Luân đao hoa tuyết tinh xảo đến từng chi tiết. Trọng lượng nhẹ như cánh hoa nhưng sắc bén vô cùng, cho phép người dùng ném liên tiếp nhiều đòn trong chớp mắt.",
    33: "Thanh đoản đao nhỏ gọn và giản dị. Thích hợp cho các sát thủ và đạo tặc ưa thích sự cơ động, mang lại tốc độ xuất chiêu cực nhanh.",
    34: "Đoản đao độc đáo với lưỡi dao uốn lượn như ngọn lửa bập bùng. Tạo ra ma sát nhiệt thiêu đốt vết thương của kẻ thù, gây sát thương hệ Hỏa.",
    35: "Đoản đao kết hợp với găng tay bảo hộ. Thiết kế thông minh cho phép người dùng vừa đỡ đòn vừa phản công tức thì trong cận chiến.",
    36: "Thanh đoản đao trong suốt tuyệt đẹp được mài từ khối pha lê ngàn năm. Khúc xạ ánh sáng tạo ra ảo ảnh quang học đánh lừa thị giác kẻ địch.",
    37: "Đoản đao chế tác từ quặng Jael, vật liệu cứng nhất được biết đến trên lục địa. Khả năng xuyên giáp tuyệt hảo, bỏ qua một phần phòng thủ của mục tiêu.",
    38: "Con dao từng được dùng trong các nghi lễ hiến tế cổ đại. Mang theo năng lượng tâm linh huyền bí, gia tăng uy lực cho các đòn tấn công ma thuật.",
    39: "Đoản đao đen nhánh phát ra ánh sáng ma mị quyến rũ. Được tẩm chất độc tê liệt chiết xuất từ loài nhện độc sâu trong hang tuyết.",
    40: "Chuôi đoản đao được đính một viên ngọc mắt quỷ bắt mắt. Giúp người dùng nhìn thấu sơ hở của đối phương và tăng gấp đôi tỉ lệ đánh chí mạng.",
    41: "Đoản đao rèn theo bí thuật viễn xứ. Cân bằng trọng lượng hoàn hảo, cho phép người dùng vung đao nhanh như bóng ma lướt qua.",
    42: "Đoản đao rèn từ quặng phát quang bí ẩn. Tỏa sáng trong đêm tối và tích lũy năng lượng để giải phóng đòn tấn công bất ngờ đầy uy lực.",
    43: "Con dao làm bếp sắc bén. Dù bề ngoài như dụng cụ nấu nướng thông thường, độ bén phi thường của nó khiến bất kỳ kẻ địch nào cũng phải khiếp sợ.",
    44: "Đoản đao đồng bộ ma lực hoàn hảo với chủ nhân. Càng chiến đấu lâu, mối liên kết càng mạnh mẽ và gia tăng chỉ số tấn công liên tục.",
    45: "Con dao cổ từng là biểu tượng sùng bái của một giáo phái ngàn năm trước. Mang sức mạnh bảo hộ linh thiêng cho người cầm lái.",
    46: "Thanh đại đao dài và mộc mạc. Vũ khí tiêu chuẩn của các cựu chiến binh dạn dày kinh nghiệm, mang lại uy lực đòn chém vô cùng nặng nề.",
    47: "Thanh kiếm nghi lễ thời cổ đại được trang trí hoa văn rực rỡ. Tuy dùng trong tế lễ, lưỡi kiếm vẫn giữ nguyên sức sát thương khủng khiếp.",
    48: "Đại đao mang hình dáng như một phiến đá nguyên khối khổng lồ. Sức nặng ngàn cân của nó nghiền nát mọi giáp trụ khi giáng xuống.",
    49: "Thanh cự kiếm uy dũng trông như tảng đá tảng vững chãi. Cho phép người dùng đứng vững như bàn thạch trước mọi đòn tấn công vũ bão.",
    50: "Đại đao được rèn hoàn toàn từ hợp kim Levissium siêu bền. Giảm đáng kể trọng lượng của cự kiếm, giúp vung những nhát chém nhanh nhẹn bất ngờ.",
    51: "Thanh kiếm cốt phát ra ánh sáng trắng mờ ảo. Được đẽo gọt từ xương của loài rồng biển cổ đại, mang ma lực biển sâu huyền bí.",
    52: "Thanh đao kỳ dị với lưỡi kiếm hình răng cưa sắc lẹm. Xé toạc mọi chướng ngại vật và gây chảy máu liên tục cho mục tiêu.",
    53: "Thanh đại đao siêu nặng với lưỡi kiếm ánh nâu đất. Tích tụ sức mạnh của đất mẹ, tạo ra chấn động rung chuyển mặt đất mỗi khi chạm đất.",
    54: "Đại đao chế tác từ kim loại khai quật từ đáy biển sâu. Chống chịu mọi sự ăn mòn và gia tăng sức mạnh khi chiến đấu trong bão tuyết.",
    55: "Thanh kiếm khổng lồ với thiết kế tối giản thực dụng. Tập trung toàn bộ vào trọng lượng và độ bền, là vũ khí ưa thích của các kiếm sĩ lực lưỡng.",
    56: "Thanh đại đao nhận được từ Freyja nhằm hỗ trợ sứ mệnh. Mang theo ý chí kiên cường bảo vệ vùng đất và đồng đội.",
    57: "Đại đao bí ẩn được rèn từ dị giới. Mang theo nguồn năng lượng không thuộc về thế giới này, phá vỡ các quy luật vật lý thông thường.",
    58: "Thanh đại kiếm làm từ chiếc xương khổng lồ của một sinh vật huyền thoại chưa xác định. Chứa đựng sức mạnh hoang dã vô tận.",
    59: "Cây quyền trượng thanh mảnh như chiếc đũa phép. Dù vẻ ngoài nhỏ bé, nó dẫn truyền ma lực mượt mà và gia tăng tốc độ niệm chú.",
    60: "Gậy phép màu trắng mang hình cán cân công lý. Thường được sử dụng bởi các học giả và pháp sư nghiên cứu ma thuật cân bằng.",
    61: "Cây gậy phép tuyệt đẹp được chạm khắc tinh xảo. Viên pha lê trên đỉnh khuếch đại uy lực của các câu thần chú ma thuật lên gấp nhiều lần.",
    62: "Gậy phép hình vầng trăng khuyết được chế tác bởi đại pháp sư. Hấp thụ ánh sáng ban đêm để hồi phục MP cho người sử dụng.",
    63: "Gậy phép được đẽo từ một nhánh của 'Cây Thế Giới'. Mang sức sống mãnh liệt của thiên nhiên, tự động hồi phục sinh lực cho người cầm.",
    64: "Cây quyền trượng lộng lẫy quý phái. Đính nhiều loại đá quý phản chiếu ma lực, tăng cường sức tấn công của các đòn phép nguyên tố.",
    65: "Gậy phép uốn lượn hình mãng xà gắn viên Spritnite đỏ rực. Chuyên dùng để dẫn truyền ma thuật lửa, thiêu rụi kẻ thù trong tích tắc.",
    66: "Gậy phép với đôi cánh thiên thần trắng muốt trên đỉnh. Biểu tượng của sự thuần khiết, gia tăng mạnh mẽ hiệu quả của các phép thuật hệ Quang.",
    67: "Quyền trượng cổ chế tác từ di vật hoàng gia ngàn năm. Lưu giữ những câu thần chú cổ xưa giúp giảm lượng MP tiêu hao khi dùng phép.",
    68: "Gậy gỗ hình chiếc búa nhỏ ngộ nghĩnh. Dù trông như đồ chơi, nó có thể phóng ra những quả cầu năng lượng ma thuật cực mạnh.",
    69: "Gậy phép làm từ mai của loài quái vật biển cổ đại. Lớp vỏ hóa thạch giúp bảo vệ pháp sư khỏi các đòn đánh ma thuật của kẻ thù.",
    70: "Ngọn thương hoàng gia truyền đời trong vương tộc. Được rèn bằng kỹ nghệ tối cao, cân bằng hoàn hảo giữa đòn đâm hiểm hóc và thế thủ kiên cố.",
    71: "Ngọn trường thương gắn tấm khiên chắn lớn ở phần chuôi. Cho phép nữ kỵ sĩ vừa lao lên tấn công vừa che chắn an toàn cho bản thân.",
    72: "Thương băng tỏa ra hàn khí buốt giá. Mũi thương nhọn hoắt đâm xuyên qua mục tiêu và đóng băng vết thương ngay tức thì.",
    73: "Ngọn đại thương sừng sững như một pháo đài di động. Tăng cường khả năng phòng thủ và chặn đứng mọi đợt xung phong của quái vật.",
    74: "Ngọn thương bạc hình xoắn ốc tuyệt mỹ. Xoáy sâu vào điểm yếu của kẻ địch khi đâm tới, gây sát thương chí mạng cực lớn.",
    75: "Ngọn thương chế tác từ sừng của quái thú độc giác cổ đại. Mũi thương cứng cáp không gì cản nổi, đâm xuyên mọi loại giáp sắt.",
    76: "Hai linh hồn song sinh được cho là đang ngự trị trong ngọn thương này. Ban cho vũ khí khả năng tấn công kép với hai luồng năng lượng đối cực.",
    77: "Ngọn thương tỏa ánh hào quang hoàng kim mờ ảo. Được ban phước bởi hoàng gia cổ xưa, gia tăng dũng khí và sức mạnh cho người mang.",
    78: "Lưỡi hái tử thần được mang bởi <NAME=CP_0007>. Sinh ra từ bóng tối vĩnh hằng, lưỡi hái gặt hái sinh lực của mọi kẻ địch cản đường.",
    79: "Lưỡi hái hắc ám mang hình đầu lâu quái thú. Nguồn tà khí tỏa ra từ lưỡi hái gieo rắc nỗi kinh hoàng và làm suy kiệt tinh thần đối thủ.",
    80: "Lưỡi hái cổ được phát hiện trong tàn tích ngàn năm. Dù thời gian trôi qua, lưỡi hái vẫn sắc bén nguyên vẹn và mang ma lực cổ đại đáng sợ.",
    81: "Lưỡi hái răng cưa phát ra ánh sáng hoàng kim ma quái. Mỗi nhát quét đều để lại những vết thương xé rách không thể hàn gắn.",
    82: "Lưỡi hái ánh bạc mờ ảo tựa sương đêm. Cho phép người dùng vung những đường hái vô hình trong bóng tối hạ gục đối thủ bất ngờ.",
    83: "Lưỡi hái tỏa ra luồng oán khí áp đảo tột cùng. Hấp thụ sinh mệnh của những kẻ bại trận để cường hóa sức mạnh cho nhát chém tiếp theo.",
    84: "Lưỡi hái khai quật từ lăng mộ của một vị vua cổ đại. Chứa đựng lời nguyền ngàn năm có thể đoạt mạng kẻ địch trong chớp mắt.",
}

for r, text in weapon_translations.items():
    ws_wep.cell(r, 3).value = text

wb.save(excel_path)
print("Updated all 79 weapon lore entries and final note entries successfully!")
