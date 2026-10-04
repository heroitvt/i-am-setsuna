import sys
sys.stdout.reconfigure(encoding='utf-8')
import os
import re
import struct
import openpyxl
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
from cryptography.hazmat.backends import default_backend

KEY = b"8xTD|EgD|b?07QDj"
IV = b"/]s@*CxLzM!9Qd%("

def encrypt_data(data):
    cipher = Cipher(algorithms.AES(KEY), modes.CBC(IV), backend=default_backend())
    encryptor = cipher.encryptor()
    return encryptor.update(data) + encryptor.finalize()

def encode_string_field(text):
    if text is None or str(text).strip() == "":
        return b""
    text_str = str(text)
    if "|" in text_str:
        pages = re.split(r'\n?\|\n?', text_str)
    else:
        pages = [text_str]
        
    num_pages = len(pages)
    buf = bytearray()
    buf.append(num_pages)
    for p in pages:
        p_bytes = p.encode('utf-16le')
        if len(p_bytes) > 255:
            raise ValueError(f"Page byte length {len(p_bytes)} > 255 for text: {repr(p[:30])}")
        buf.append(len(p_bytes))
        buf.extend(p_bytes)
    return bytes(buf)

def main():
    target_text = "Ừm... Đúng như lời đồn.\nMột lính đánh thuê bẩm sinh."
    print("Target text for Row 6 (ID 3):")
    print(target_text)

    # 1. Update Excel
    excel_path = r"d:\Viet Hoa Game\I_am_Setsuna-parameter-Steam.xlsx"
    print(f"\n1. Updating Excel file: {excel_path}")
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["ScenarioMessageData_Chapter_1"]
    
    # Check row 6
    id_val = ws.cell(6, 1).value
    old_vn = ws.cell(6, 3).value
    print(f"Current Row 6 ID: {id_val}")
    print(f"Current Row 6 Vietnamese: {repr(old_vn)}")
    
    ws.cell(6, 3).value = target_text
    wb.save(excel_path)
    print("Successfully updated and saved Excel file!")

    # 2. Update ScenarioMessageData_Chapter_1.dec
    dec_path = r"d:\Viet Hoa Game\I am Setsuna\SETSUNA_Data\StreamingAssets\data\parameter\ScenarioMessageData_Chapter_1.dec"
    bak_path = r"d:\Viet Hoa Game\temp_scripts\ScenarioMessageData_Chapter_1.dec.bak_row6"
    
    with open(dec_path, "rb") as f:
        data = f.read()
        
    with open(bak_path, "wb") as f:
        f.write(data)
    print(f"\n2. Backed up current .dec to {bak_path}")

    new_data = bytearray()
    pos = 0
    records_count = 0
    rid3_updated = False

    while pos < len(data):
        if pos + 16 > len(data):
            break
        magic, rid, l1, l2, l3, rlen, npcid = struct.unpack("<ihhhhHh", data[pos:pos+16])
        rec_data = data[pos:pos+rlen]
        
        if rid == 3:
            cur = 16
            n_len = rec_data[cur]; cur += 1
            npc_bytes = rec_data[cur:cur+n_len]; cur += n_len
            
            q_len = rec_data[cur]; cur += 1
            q_bytes = rec_data[cur:cur+q_len]; cur += q_len
            
            s1_bytes = rec_data[cur:cur+l1]; cur += l1
            old_s2_bytes = rec_data[cur:cur+l2]; cur += l2
            s3_bytes = rec_data[cur:cur+l3]; cur += l3

            # Encode new s2
            new_s2_bytes = encode_string_field(target_text)
            new_l2 = len(new_s2_bytes)
            new_rlen = 16 + 1 + n_len + 1 + q_len + l1 + new_l2 + l3
            
            # Pack new header
            new_hdr = struct.pack("<ihhhhHh", magic, rid, l1, new_l2, l3, new_rlen, npcid)
            
            new_data.extend(new_hdr)
            new_data.append(n_len)
            new_data.extend(npc_bytes)
            new_data.append(q_len)
            new_data.extend(q_bytes)
            new_data.extend(s1_bytes)
            new_data.extend(new_s2_bytes)
            new_data.extend(s3_bytes)
            
            rid3_updated = True
            print(f"Record RID 3 updated: old_l2={l2} -> new_l2={new_l2}, old_rlen={rlen} -> new_rlen={new_rlen}")
        else:
            new_data.extend(rec_data)
            
        records_count += 1
        pos += rlen

    print(f"Processed {records_count} records. RID 3 updated: {rid3_updated}")
    assert rid3_updated, "RID 3 was not found!"
    assert records_count == 1403, f"Expected 1403 records, got {records_count}"

    # Write new .dec
    with open(dec_path, "wb") as f:
        f.write(new_data)
    print(f"Written new .dec ({len(new_data)} bytes) to {dec_path}")

    # 3. Encrypt and write loose file
    enc_loose_path = r"d:\Viet Hoa Game\I am Setsuna\SETSUNA_Data\StreamingAssets\data\parameter\ScenarioMessageData_Chapter_1"
    packed_temp_path = r"d:\Viet Hoa Game\temp_scripts\ScenarioMessageData_Chapter_1.packed"

    data_to_enc = bytearray(new_data)
    pad_len = (16 - (len(data_to_enc) % 16)) % 16
    if pad_len > 0:
        data_to_enc.extend(b"\x00" * pad_len)
        
    encrypted = encrypt_data(bytes(data_to_enc))
    with open(enc_loose_path, "wb") as f:
        f.write(encrypted)
    with open(packed_temp_path, "wb") as f:
        f.write(encrypted)
    print(f"Encrypted and saved ({len(encrypted)} bytes) to {enc_loose_path} and {packed_temp_path}")

if __name__ == "__main__":
    main()
