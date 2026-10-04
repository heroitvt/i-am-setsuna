# -*- coding: utf-8 -*-
"""Apply all 5 translation batches to I_am_Setsuna-parameter-Steam.xlsx and ScenarioMessageData_Chapter_1.json"""

import json
import shutil
import openpyxl
import re

from batch1_trans import batch1_data
from batch2_trans import batch2_data
from batch3_trans import batch3_data
from batch4_trans import batch4_data
from batch5_trans import batch5_data

excel_path = 'd:/Viet Hoa Game/I_am_Setsuna-parameter-Steam.xlsx'
backup_path = 'd:/Viet Hoa Game/I_am_Setsuna-parameter-Steam.xlsx.bak'
json_path = 'translation_export/ScenarioMessageData_Chapter_1.json'
missing_json_path = 'chapter_1_missing.json'

print('Step 1: Merging all batches...')
all_trans = {}
all_trans.update(batch1_data)
all_trans.update(batch2_data)
all_trans.update(batch3_data)
all_trans.update(batch4_data)
all_trans.update(batch5_data)

print(f'Total translated lines: {len(all_trans)}')
assert len(all_trans) == 897, f'Expected 897 lines, got {len(all_trans)}'

for r in range(509, 1406):
    assert r in all_trans, f'Missing row {r}'

print('Step 2: Backing up Excel file...')
shutil.copy2(excel_path, backup_path)
print(f'Backup saved to {backup_path}')

print('Step 3: Updating Excel workbook...')
wb = openpyxl.load_workbook(excel_path)
ws = wb['ScenarioMessageData_Chapter_1']

excel_updated = 0
for r in range(509, 1406):
    vi_text = all_trans[r]
    ws.cell(row=r, column=3).value = vi_text
    excel_updated += 1

print(f'Updated {excel_updated} rows in Excel.')
wb.save(excel_path)
print('Excel workbook saved successfully.')

print('Step 4: Updating JSON file...')
with open(json_path, 'r', encoding='utf-8') as f:
    json_data = json.load(f)

with open(missing_json_path, 'r', encoding='utf-8') as f:
    missing_data = json.load(f)

# missing_data has row numbers
missing_by_id = {item['id']: item for item in missing_data}

json_updated = 0
for item in json_data:
    item_id = item['id']
    if item_id in missing_by_id:
        row_num = missing_by_id[item_id]['row']
        if row_num in all_trans:
            item['vietnamese'] = all_trans[row_num]
            json_updated += 1

print(f'Updated {json_updated} items in JSON.')
assert json_updated == 897, f'Expected 897 updated JSON items, got {json_updated}'

with open(json_path, 'w', encoding='utf-8') as f:
    json.dump(json_data, f, ensure_ascii=False, indent=2)
print('JSON file saved successfully.')

print('Step 5: Verifying final state in Excel and JSON...')
wb_check = openpyxl.load_workbook(excel_path, data_only=True)
ws_check = wb_check['ScenarioMessageData_Chapter_1']

total_dialogue = ws_check.max_row - 1 # excluding header row 1
untranslated_in_excel = []
tag_mismatches = []
tag_re = re.compile(r'<[^>]+>')

for r in range(2, ws_check.max_row + 1):
    if r == 2:
        continue # row 2 is system parameter '1'
    eng = ws_check.cell(row=r, column=2).value
    vi = ws_check.cell(row=r, column=3).value
    if vi is None or str(vi).strip() == '':
        untranslated_in_excel.append((r, eng))
    else:
        eng_tags = sorted(tag_re.findall(str(eng or '')))
        vi_tags = sorted(tag_re.findall(str(vi or '')))
        if eng_tags != vi_tags:
            tag_mismatches.append((r, eng_tags, vi_tags))

print(f'Verification Results:')
print(f'Total rows in sheet Chapter 1: {ws_check.max_row}')
print(f'Total dialogue rows checked: {ws_check.max_row - 2}')
print(f'Untranslated rows remaining in Excel: {len(untranslated_in_excel)}')
print(f'Tag mismatches found: {len(tag_mismatches)}')

with open(json_path, 'r', encoding='utf-8') as f:
    json_check = json.load(f)

json_untrans = [x for x in json_check if not x.get('vietnamese')]
print(f'Untranslated items remaining in JSON: {len(json_untrans)}')

if len(untranslated_in_excel) == 0 and len(tag_mismatches) == 0 and len(json_untrans) == 0:
    print('SUCCESS: Chapter 1 is 100% translated and fully verified in both Excel and JSON!')
else:
    print('WARNING: Some items still need attention.')
