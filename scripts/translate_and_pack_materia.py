import openpyxl
import struct
import os
import re
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
from cryptography.hazmat.backends import default_backend

KEY = b"8xTD|EgD|b?07QDj"
IV = b"/]s@*CxLzM!9Qd%("

def decrypt_data(data):
    cipher = Cipher(algorithms.AES(KEY), modes.CBC(IV), backend=default_backend())
    decryptor = cipher.decryptor()
    return decryptor.update(data) + decryptor.finalize()

def encrypt_data(data):
    cipher = Cipher(algorithms.AES(KEY), modes.CBC(IV), backend=default_backend())
    encryptor = cipher.encryptor()
    return encryptor.update(data) + encryptor.finalize()

def main():
    excel_path = r"D:\Viet Hoa Game\I_am_Setsuna-parameter-Steam.xlsx"
    wb = openpyxl.load_workbook(excel_path)
    
    # 1. Build skill map from 7 character skill sheets
    skill_map = {}
    for sname in ["PlayerSkillDataMessage", "SetsunaSkillDataMessage", "TsukushiSkillDataMessage", "YomiSkillDataMessage", "KishilSkillDataMessage", "SionSkillDataMessage", "GrimreaperSkillDataMessage"]:
        ws = wb[sname]
        r = 6
        while r <= ws.max_row:
            en_name = ws.cell(r, 2).value
            vn_name = ws.cell(r, 3).value
            en_desc = ws.cell(r+1, 2).value
            vn_desc = ws.cell(r+1, 3).value
            if en_name and vn_name:
                skill_map[str(en_name).strip()] = {
                    "name": str(vn_name).strip(),
                    "desc": str(vn_desc).strip() if vn_desc else ""
                }
            r += 2

    # Support Spritnite names and descriptions dictionary
    support_names = {
        "Lifeforce Logic": ("Logic Sinh Mệnh", "Dùng ma lực bảo tồn sinh mệnh, tự động hồi một lượng nhỏ HP khi bắt đầu mỗi lượt."),
        "Magical Pulse": ("Mạch Ma Lực", "Dùng ma lực mở rộng tâm trí, tự động hồi một lượng nhỏ MP khi bắt đầu mỗi lượt."),
        "Physical Pride": ("Lòng Kiêu Hãnh", "Đánh thức tiềm năng thể chất, tăng vĩnh viễn chỉ số Tấn Công vật lý."),
        "Wise Testimony": ("Minh Triết Tối Cao", "Đánh thức tiềm năng pháp thuật, tăng vĩnh viễn chỉ số Tấn Công phép thuật."),
        "Iron Vow": ("Lời Thề Thép", "Đánh thức sức mạnh phòng hộ, tăng vĩnh viễn chỉ số Phòng Thủ vật lý."),
        "Soul Prayer": ("Lời Cầu Nguyện Linh Hồn", "Đánh thức sức mạnh kháng phép, tăng vĩnh viễn chỉ số Phòng Thủ phép thuật."),
        "Rising Spirit": ("Ý Chí Bất Khuất", "Tăng cường độ nhạy bén của các giác quan, tăng chỉ số Đòn Chí Mạng."),
        "All-Seeing Eye": ("Mắt Thần Toàn Tri", "Tăng cường độ tập trung tuyệt đối, tăng độ chính xác của mọi đòn đánh."),
        "Ultimate Truth": ("Chân Lý Tối Thượng", "Tăng cường độ linh hoạt phản xạ, tăng tỉ lệ Né Tránh đòn đánh của kẻ địch."),
        "Universal Insight": ("Tuệ Nhãn Vũ Trụ", "Tăng cường khả năng nắm bắt sơ hở của kẻ địch, tăng tỉ lệ Đánh Trúng Điểm Yếu."),
        "Cursed Helix": ("Vòng Xoáy Nguyền Rủa", "Tăng cường năng lượng nguyền rủa, tăng tỉ lệ áp đặt các trạng thái bất lợi lên kẻ địch."),
        "Commanding Wave": ("Sóng Lệnh Thời Gian", "Can thiệp vào dòng chảy thời gian, tăng tốc độ nạp thanh đo ATB."),
        "Wailing Wind": ("Gió Rít Gào Thét", "Can thiệp vào không gian chiến trường, tăng tốc độ di chuyển của nhân vật."),
        "Incandescent Instant": ("Khoảnh Khắc Bùng Cháy", "Can thiệp vào thời gian niệm chú, giảm thời gian chờ khi sử dụng kỹ năng."),
        "Fated Memory": ("Ký Ức Định Mệnh", "Can thiệp vào thời gian, nạp sẵn thanh ATB khi kích hoạt chế độ Xung Lực."),
        "Valiant Poem": ("Bài Thơ Quả Cảm", "Can thiệp vào ý chí chiến đấu, nạp sẵn thanh SP khi nhân vật tử trận được hồi sinh."),
        "Guiding Power": ("Sức Mạnh Dẫn Lối", "Tăng tốc độ tích lũy thanh đo SP lên mức tối đa."),
        "Destined Cycle": ("Vòng Luân Hồi Định Mệnh", "Tăng lượng điểm SP tích lũy nhận được khi nhân vật bị tấn công."),
        "Heavenly Miracle": ("Phép Màu Thiên Giới", "Tăng lượng điểm SP tích lũy nhận được khi đồng minh bị tiêu diệt."),
        "Rising Reverie": ("Mộng Tưởng Thăng Hoa", "Tăng lượng điểm SP tích lũy nhận được khi nhân vật hạ gục kẻ địch."),
        "Soul Rage": ("Cuồng Nộ Linh Hồn", "Tăng uy lực đòn đánh và kỹ năng khi kích hoạt ở chế độ Xung Lực."),
        "Binding Promise": ("Lời Hứa Ràng Buộc", "Kéo dài thời gian duy trì của các hiệu ứng kích hoạt từ chế độ Xung Lực."),
        "Joyful Song": ("Khúc Ca Hoan Hỉ", "Tăng tỉ lệ xuất hiện của hiệu ứng toàn đội Điểm Dị Thường (Singularity)."),
        "Attack Bit": ("Cầu Hộ Thân Tấn Công", "Tạo ra một quả cầu vô hình bay quanh, tự động tung đòn tấn công phụ khi đánh thường."),
        "Fire Bit": ("Cầu Hộ Thân Hệ Hỏa", "Tự động tung đòn tấn công phụ hệ Hỏa khi thực hiện hành động."),
        "Water Bit": ("Cầu Hộ Thân Hệ Băng", "Tự động tung đòn tấn công phụ hệ Băng/Thủy khi thực hiện hành động."),
        "Light Bit": ("Cầu Hộ Thân Hệ Quang", "Tự động tung đòn tấn công phụ hệ Quang khi thực hiện hành động."),
        "Shadow Bit": ("Cầu Hộ Thân Hệ Hắc Ám", "Tự động tung đòn tấn công phụ hệ Hắc Ám khi thực hiện hành động."),
        "Time Bit": ("Cầu Hộ Thân Hệ Thời Gian", "Tự động tung đòn tấn công phụ hệ Thời Gian khi thực hiện hành động."),
        "Fire Shield": ("Khiên Hộ Thể Hệ Hỏa", "Tạo màn chắn ma lực tăng mạnh khả năng kháng sát thương hệ Hỏa."),
        "Water Shield": ("Khiên Hộ Thể Hệ Băng", "Tạo màn chắn ma lực tăng mạnh khả năng kháng sát thương hệ Băng/Thủy."),
        "Light Shield": ("Khiên Hộ Thể Hệ Quang", "Tạo màn chắn ma lực tăng mạnh khả năng kháng sát thương hệ Quang."),
        "Shadow Shield": ("Khiên Hộ Thể Hệ Hắc Ám", "Tạo màn chắn ma lực tăng mạnh khả năng kháng sát thương hệ Hắc Ám."),
        "Time Shield": ("Khiên Hộ Thể Hệ Thời Gian", "Tạo màn chắn ma lực tăng mạnh khả năng kháng sát thương hệ Thời Gian."),
        "Sap Bit": ("Cầu Hộ Thân Ăn Mòn", "Tự động tung đòn tấn công có xác suất gây hiệu ứng Rút Máu (Sap) lên kẻ địch."),
        "Paralysis Bit": ("Cầu Hộ Thân Tê Liệt", "Tự động tung đòn tấn công có xác suất gây hiệu ứng Tê Liệt (Paralysis)."),
        "Confusion Bit": ("Cầu Hộ Thân Hỗn Loạn", "Tự động tung đòn tấn công có xác suất gây hiệu ứng Hỗn Loạn (Confusion)."),
        "Stone Bit": ("Cầu Hộ Thân Hóa Đá", "Tự động tung đòn tấn công có xác suất gây hiệu ứng Hóa Đá (Stone)."),
        "Freeze Bit": ("Cầu Hộ Thân Đóng Băng", "Tự động tung đòn tấn công có xác suất gây hiệu ứng Đóng Băng (Freeze)."),
        "Stop Bit": ("Cầu Hộ Thân Ngưng Đọng", "Tự động tung đòn tấn công có xác suất gây hiệu ứng Ngưng Đọng Thời Gian (Stop)."),
        "Stun Bit": ("Cầu Hộ Thân Choáng Váng", "Tự động tung đòn tấn công có xác suất gây hiệu ứng Choáng Váng (Stun)."),
        "Death Bit": ("Cầu Hộ Thân Tử Vong", "Tự động tung đòn tấn công có xác suất gây hiệu ứng Đoạt Mạng Tức Thì (Death)."),
        "Sap Shield": ("Khiên Kháng Rút Máu", "Tăng khả năng miễn nhiễm hoàn toàn trước trạng thái Rút Máu (Sap)."),
        "Paralysis Shield": ("Khiên Kháng Tê Liệt", "Tăng khả năng miễn nhiễm hoàn toàn trước trạng thái Tê Liệt (Paralysis)."),
        "Confusion Shield": ("Khiên Kháng Hỗn Loạn", "Tăng khả năng miễn nhiễm hoàn toàn trước trạng thái Hỗn Loạn (Confusion)."),
        "Stone Shield": ("Khiên Kháng Hóa Đá", "Tăng khả năng miễn nhiễm hoàn toàn trước trạng thái Hóa Đá (Stone)."),
        "Freeze Shield": ("Khiên Kháng Đóng Băng", "Tăng khả năng miễn nhiễm hoàn toàn trước trạng thái Đóng Băng (Freeze)."),
        "Stop Shield": ("Khiên Kháng Ngưng Đọng", "Tăng khả năng miễn nhiễm hoàn toàn trước trạng thái Ngưng Đọng (Stop)."),
        "Lock Shield": ("Khiên Kháng Khóa Phép", "Tăng khả năng miễn nhiễm hoàn toàn trước trạng thái Khóa Kỹ Năng (Lock)."),
        "Stun Shield": ("Khiên Kháng Choáng", "Tăng khả năng miễn nhiễm hoàn toàn trước trạng thái Choáng Váng (Stun)."),
        "Death Shield": ("Khiên Kháng Tử Vong", "Tăng khả năng miễn nhiễm hoàn toàn trước các đòn đánh Đoạt Mạng Tức Thì (Death)."),
        "Resonance": ("Cộng Hưởng", "Kích thích ma lực bản thân, tăng mạnh lượng HP hồi phục từ các kỹ năng trị thương."),
        "Vinculum": ("Liên Kết Bất Diệt", "Kích thích ma lực bản thân, mở rộng phạm vi tác dụng của các kỹ năng hỗ trợ."),
        "Fatus": ("Định Mệnh", "Kích thích ma lực bản thân, kéo dài thời gian duy trì của các bùa lợi hỗ trợ."),
        "Kraftwerk": ("Cường Lực", "Kích thích ma lực bản thân, giảm một nửa lượng MP tiêu hao của mọi kỹ năng."),
        "Calor": ("Nhiệt Lượng", "Kích thích ma lực bản thân, tăng tốc độ nạp thanh ATB sau khi thực hiện hành động."),
        "Oratio": ("Nguyện Cầu", "Kích thích ma lực bản thân, tăng tỉ lệ kích hoạt Biến Chuyển Flux khi dùng chế độ Xung Lực."),
        "Esperanza": ("Hi Vọng", "Kích thích ma lực bản thân, nạp đầy thanh ATB ngay khi bắt đầu trận chiến."),
        "Megalith": ("Cự Thạch", "Chuyển hóa ma lực thành lá chắn, miễn nhiễm với mọi sát thương vật lý khi HP đầy 100%."),
        "Enigma": ("Bí Ẩn", "Chuyển hóa ma lực thành lá chắn, miễn nhiễm với mọi sát thương phép thuật khi HP đầy 100%."),
        "Sigtyr": ("Thần Khí", "Hấp thụ ma lực kẻ địch, chuyển hóa 10% sát thương vật lý gây ra thành HP cho bản thân."),
        "Gagnrath": ("Thần Lực", "Hấp thụ ma lực kẻ địch, chuyển hóa 5% sát thương phép thuật gây ra thành MP cho bản thân."),
        "Bait": ("Mồi Nhử", "Dùng ma lực thu hút sự chú ý của kẻ địch, khiến chúng luôn nhắm vào bản thân."),
        "Cloak": ("Áo Choàng Ẩn Thân", "Dùng ma lực che giấu hoàn toàn sự hiện diện, khiến kẻ địch không thể nhắm vào bản thân."),
        "Returner": ("Phản Kích", "Chuyển hóa đòn tấn công của kẻ địch thành phản đòn vật lý chớp nhoáng."),
        "Avenger": ("Báo Thù", "Chuyển hóa đòn tấn công của kẻ địch thành phản đòn ma thuật uy lực."),
        "Inner Calm": ("Tĩnh Tâm", "Thay đổi quy luật tự nhiên, giữ nguyên thanh ATB khi bị trúng đòn tấn công của kẻ địch."),
        "Perfect Composure": ("Điềm Tĩnh Tuyệt Đối", "Thay đổi quy luật tự nhiên, giữ nguyên thanh SP khi bị trúng đòn tấn công của kẻ địch."),
        "Myriad Power": ("Vạn Lực", "Khi tấn công gây sát thương chí mạng, hồi phục lại 20% lượng MP tiêu hao."),
        "Unstoppable Force": ("Bất Khả Cản Phá", "Khi dùng kỹ năng hoặc liên chiêu, đòn đánh không thể bị kẻ địch ngắt quãng."),
        "Eternal Recurrence": ("Vĩnh Cửu Luân Hồi", "Khi dùng kỹ năng hoặc liên chiêu, có tỉ lệ không bị tiêu hao MP."),
        "Kaleidoscopic Shift": ("Biến Chuyển Muôn Màu", "Khi dùng kỹ năng hoặc liên chiêu, đòn đánh sẽ mang toàn bộ các thuộc tính nguyên tố."),
        "Earth Shaker": ("Địa Chấn", "Đòn đánh vật lý của nhân vật sẽ bỏ qua 30% chỉ số Phòng Thủ của kẻ địch."),
        "Combat Instinct": ("Bản Năng Chiến Đấu", "Đòn đánh phép thuật của nhân vật sẽ bỏ qua 30% chỉ số Kháng Phép của kẻ địch."),
        "Final Comeback": ("Lội Ngược Dòng", "Khi HP của nhân vật tụt xuống dưới 30%, tăng gấp đôi chỉ số Tấn Công và Phòng Thủ."),
        "Spiritual Harmony": ("Hài Hòa Linh Hồn", "Tăng lượng HP và MP hồi phục nhận được từ toàn bộ các nguồn bổ trợ."),
        "Transcendent Mind": ("Tâm Trí Siêu Việt", "Giảm 30% lượng sát thương phép thuật phải gánh chịu từ mọi nguồn."),
        "Radiant Heart": ("Trái Tim Tỏa Sáng", "Giảm 30% lượng sát thương vật lý phải gánh chịu từ mọi nguồn."),
        "Perpetual Fortitude": ("Kiên Cường Vĩnh Cửu", "Bất kỳ đòn đánh nào có thể khiến nhân vật tử trận sẽ giữ lại cho nhân vật đúng 1 HP (1 lần/trận)."),
        "Harmonic Balance": ("Cân Bằng Hài Hòa", "Mỗi đòn đánh thành công sẽ hồi phục một lượng nhỏ HP và MP cho toàn đội."),
        "Selfless Devotion": ("Cống Hiến Vị Tha", "Tăng sức mạnh của mọi hành động dựa trên số lần kích hoạt chế độ Xung Lực trong trận."),
        "Ceaseless Shift": ("Chuyển Dịch Không Ngừng", "Tăng sức mạnh của mọi hành động dựa trên số lần kích hoạt Biến Chuyển Flux trong trận."),
        "Dynamic Spirit": ("Tinh Thần Năng Động", "Tăng tốc độ tích lũy thanh ATB cho toàn bộ các thành viên trong đội."),
        "Infinite Ingenuity": ("Mưu Trí Vô Hạn", "Tăng lượng điểm kinh nghiệm (EXP) và tiền vàng (Gold) nhận được sau mỗi trận đấu."),
        "Mental Emancipation": ("Giải Phóng Tâm Trí", "Tăng tỉ lệ rơi ra các nguyên liệu và vật phẩm quý hiếm từ quái vật."),
        "Eternal Void": ("Hư Vô Vĩnh Hằng", "Miễn nhiễm với toàn bộ các trạng thái bất lợi thông thường."),
        "Berserker": ("Cuồng Chiến Binh", "Trong trận chiến, có xác suất rơi vào trạng thái Cuồng Nộ, tăng cực hạn chỉ số Tấn Công."),
        "Unlimited": ("Vô Hạn", "Trong trận chiến, có xác suất rơi vào trạng thái Vô Hạn, toàn bộ kỹ năng không tốn MP."),
        "Bifrost": ("Cầu Cầu Vồng", "Trong trận chiến, có xác suất được ban toàn bộ các bùa lợi hỗ trợ trong 3 lượt."),
        "Monad": ("Đơn Nguyên", "Trong trận chiến, có xác suất làm đầy 100% thanh ATB và SP của toàn đội ngay lập tức."),
        "Alfadir": ("Thần Phụ", "Trong trận chiến, có xác suất hồi sinh toàn bộ đồng minh đã tử trận với 100% HP."),
        "Superbia": ("Kiêu Hãnh Thần Thánh", "Trong trận chiến, có xác suất tăng gấp đôi toàn bộ chỉ số của bản thân."),
        "Laplace": ("Ma Trận Laplace", "Trong trận chiến, có xác suất khiến toàn bộ đòn đánh của nhân vật chắc chắn chí mạng."),
        "Yggdrasil": ("Cây Thế Giới", "Trong trận chiến, có xác suất hồi phục toàn bộ 100% HP và MP cho toàn đội."),
        "Creare": ("Sáng Thế", "Trong trận chiến, có xác suất lập tức kích hoạt hiệu ứng Điểm Dị Thường."),
        "Overload": ("Quá Tải Năng Lượng", "Trong trận chiến, có xác suất giải phóng đòn tấn công ma thuật cực mạnh lên toàn bộ kẻ địch."),
        "Omniverse": ("Đa Vũ Trụ", "Khi dùng kỹ năng ở chế độ Xung Lực, đòn đánh không thể bị kẻ địch phản lại."),
        "Reginleif": ("Thần Nữ Chiến Trận", "Khi dùng kỹ năng ở chế độ Xung Lực, sát thương gây ra chuyển thành sát thương đặc biệt bỏ qua phòng thủ."),
        "Clairvoyance": ("Thấu Thị", "Khi dùng kỹ năng ở chế độ Xung Lực, đòn đánh sẽ mang toàn bộ các thuộc tính nguyên tố."),
        "Gratia": ("Ân Huệ", "Khi dùng liên chiêu combo ở chế độ Xung Lực, thanh ATB sẽ được nạp đầy ngay sau khi tung chiêu."),
        "Alchemia": ("Giả Kim Thuật", "Khi dùng vật phẩm ở chế độ Xung Lực, vật phẩm đó trong túi đồ sẽ không bị tiêu hao."),
        "Paries": ("Tường Thành", "Kích hoạt chế độ Xung Lực khi bị tấn công sẽ giảm một nửa lượng sát thương phải nhận."),
        "Radgrid": ("Khiên Thánh", "Kích hoạt chế độ Xung Lực khi bị tấn công có trạng thái bất lợi sẽ lập tức hóa giải trạng thái đó."),
        "Reverse": ("Nghịch Chuyển", "Kích hoạt chế độ Xung Lực khi bị tấn công sẽ phản lại một phần sát thương cho kẻ địch."),
        "Tactical Combo": ("Liên Hoàn Chiến Thuật", "Nếu thực hiện liên tiếp các kỹ năng cấu thành liên chiêu, dùng kỹ năng cuối ở chế độ Xung Lực sẽ tự động kích hoạt combo."),
        "Revival Counter": ("Phản Kích Hồi Sinh", "Kích hoạt chế độ Xung Lực khi bị tấn công trong trạng thái HP dưới 20% sẽ lập tức tung đòn phản công cực mạnh.")
    }

    # 2. Update Excel sheet MateriaMessage
    ws_mat = wb["MateriaMessage"]
    materia_items = []

    for r in range(6, ws_mat.max_row + 1, 4):
        en_name = str(ws_mat.cell(r, 2).value or "").strip()
        en_desc = str(ws_mat.cell(r+1, 2).value or "").strip()
        
        vn_name = ""
        vn_desc = ""

        if en_name in skill_map:
            vn_name = skill_map[en_name]["name"]
            vn_desc = f"Trang bị để sử dụng kỹ năng {vn_name}.\r\n{skill_map[en_name]['desc']}"
        elif en_name in support_names:
            vn_name, s_desc = support_names[en_name]
            vn_desc = f"Kích hoạt {vn_name} trong chiến đấu.\r\n{s_desc}"
        else:
            vn_name = en_name
            vn_desc = en_desc

        # Update in Excel
        ws_mat.cell(r, 3, vn_name)
        ws_mat.cell(r+1, 3, vn_desc)
        materia_items.append((vn_name, vn_desc))

    print(f"Updated {len(materia_items)} Materia entries in Excel!")
    wb.save(excel_path)
    print("Saved Excel successfully!")

    # 3. Re-pack binary MateriaMessage
    bin_path = r"D:\Viet Hoa Game\I am Setsuna\SETSUNA_Data\StreamingAssets\data\parameter\MateriaMessage"
    with open(bin_path, "rb") as f:
        encrypted_raw = f.read()

    dec = bytearray(decrypt_data(encrypted_raw))
    num_records = struct.unpack("<h", dec[2:4])[0]
    print(f"Packing MateriaMessage: {len(materia_items)} items, {num_records} binary records.")

    offset = 4
    new_data = bytearray()
    new_data.extend(dec[:4])

    for i in range(num_records):
        # Name
        jp_n_len = struct.unpack("<h", dec[offset:offset+2])[0]; offset += 2
        jp_n_bytes = dec[offset:offset+jp_n_len]; offset += jp_n_len

        en_n_len = struct.unpack("<h", dec[offset:offset+2])[0]; offset += 2
        en_n_bytes = dec[offset:offset+en_n_len]; offset += en_n_len

        fr_n_len = struct.unpack("<h", dec[offset:offset+2])[0]; offset += 2
        fr_n_bytes = dec[offset:offset+fr_n_len]; offset += fr_n_len

        # Desc
        jp_d_len = struct.unpack("<h", dec[offset:offset+2])[0]; offset += 2
        jp_d_bytes = dec[offset:offset+jp_d_len]; offset += jp_d_len

        en_d_len = struct.unpack("<h", dec[offset:offset+2])[0]; offset += 2
        en_d_bytes = dec[offset:offset+en_d_len]; offset += en_d_len

        fr_d_len = struct.unpack("<h", dec[offset:offset+2])[0]; offset += 2
        fr_d_bytes = dec[offset:offset+fr_d_len]; offset += fr_d_len

        if i < len(materia_items):
            vn_n, vn_d = materia_items[i]
            new_en_n_bytes = vn_n.encode("utf-16le")
            new_en_n_len = len(new_en_n_bytes)
            new_en_d_bytes = vn_d.encode("utf-16le")
            new_en_d_len = len(new_en_d_bytes)
        else:
            new_en_n_bytes = en_n_bytes
            new_en_n_len = en_n_len
            new_en_d_bytes = en_d_bytes
            new_en_d_len = en_d_len

        # Write
        new_data.extend(struct.pack("<h", jp_n_len))
        new_data.extend(jp_n_bytes)
        new_data.extend(struct.pack("<h", new_en_n_len))
        new_data.extend(new_en_n_bytes)
        new_data.extend(struct.pack("<h", fr_n_len))
        new_data.extend(fr_n_bytes)

        new_data.extend(struct.pack("<h", jp_d_len))
        new_data.extend(jp_d_bytes)
        new_data.extend(struct.pack("<h", new_en_d_len))
        new_data.extend(new_en_d_bytes)
        new_data.extend(struct.pack("<h", fr_d_len))
        new_data.extend(fr_d_bytes)

    with open(bin_path + ".dec", "wb") as f:
        f.write(new_data)

    pad_len = 16 - (len(new_data) % 16)
    if pad_len < 16:
        new_data.extend(b"\x00" * pad_len)

    encrypted = encrypt_data(bytes(new_data))
    with open(bin_path, "wb") as f:
        f.write(encrypted)

    patch_dest = os.path.join(r"D:\Viet Hoa Game\Setsuna_VietHoa_Patch\SETSUNA_Data\StreamingAssets\data\parameter\MateriaMessage")
    with open(patch_dest, "wb") as f:
        f.write(encrypted)

    print(f"Successfully packed and encrypted MateriaMessage ({len(encrypted)} bytes)!")

if __name__ == "__main__":
    main()
