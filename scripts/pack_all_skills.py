import os
import struct
import openpyxl
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
from cryptography.hazmat.backends import default_backend

KEY = b"8xTD|EgD|b?07QDj"
IV = b"/]s@*CxLzM!9Qd%("

def decrypt_data(data):
    cipher = Cipher(algorithms.AES(KEY), modes.CBC(IV), backend=default_backend())
    decryptor = cipher.decryptor()
    return decryptor.update(data) + decryptor.finalize()

def encrypt_data(data):
    cipher = Cipher(algorithms.AES(KEY), modes.CBC(IV), backend=default_backend())
    encryptor = cipher.encryptor()
    return encryptor.update(data) + encryptor.finalize()

def pack_skill_file(excel_sheet_name, binary_file_name, wb):
    ws = wb[excel_sheet_name]
    pairs = []
    r = 6
    while r <= ws.max_row:
        name_en = ws.cell(r, 2).value
        name_vn = ws.cell(r, 3).value
        desc_en = ws.cell(r+1, 2).value if r+1 <= ws.max_row else ""
        desc_vn = ws.cell(r+1, 3).value if r+1 <= ws.max_row else ""
        if name_en or name_vn:
            pairs.append({
                "name": str(name_vn if name_vn and str(name_vn).strip() else (name_en or "")).strip(),
                "desc": str(desc_vn if desc_vn and str(desc_vn).strip() else (desc_en or "")).strip()
            })
        r += 2

    bin_path = os.path.join(r"D:\Viet Hoa Game\I am Setsuna\SETSUNA_Data\StreamingAssets\data\parameter", binary_file_name)
    with open(bin_path, "rb") as f:
        encrypted_raw = f.read()

    dec = bytearray(decrypt_data(encrypted_raw))
    num_records = struct.unpack("<h", dec[2:4])[0]
    print(f"Packing {binary_file_name}: {len(pairs)} Excel pairs, {num_records} binary records.")

    offset = 4
    new_data = bytearray()
    new_data.extend(dec[:4])

    for i in range(num_records):
        # 1. Read original names (3 langs)
        jp_n_len = struct.unpack("<h", dec[offset:offset+2])[0]; offset += 2
        jp_n_bytes = dec[offset:offset+jp_n_len]; offset += jp_n_len

        en_n_len = struct.unpack("<h", dec[offset:offset+2])[0]; offset += 2
        en_n_bytes = dec[offset:offset+en_n_len]; offset += en_n_len

        fr_n_len = struct.unpack("<h", dec[offset:offset+2])[0]; offset += 2
        fr_n_bytes = dec[offset:offset+fr_n_len]; offset += fr_n_len

        # 2. Read original descs (3 langs)
        jp_d_len = struct.unpack("<h", dec[offset:offset+2])[0]; offset += 2
        jp_d_bytes = dec[offset:offset+jp_d_len]; offset += jp_d_len

        en_d_len = struct.unpack("<h", dec[offset:offset+2])[0]; offset += 2
        en_d_bytes = dec[offset:offset+en_d_len]; offset += en_d_len

        fr_d_len = struct.unpack("<h", dec[offset:offset+2])[0]; offset += 2
        fr_d_bytes = dec[offset:offset+fr_d_len]; offset += fr_d_len

        # If we have translation in pairs:
        if i < len(pairs):
            vn_name = pairs[i]["name"]
            vn_desc = pairs[i]["desc"]
            # Convert \n to \r\n for standard Unity dialogue box if needed
            vn_desc = vn_desc.replace("\r\n", "\n").replace("\n", "\r\n")

            new_en_n_bytes = vn_name.encode("utf-16le")
            new_en_n_len = len(new_en_n_bytes)

            new_en_d_bytes = vn_desc.encode("utf-16le")
            new_en_d_len = len(new_en_d_bytes)
        else:
            new_en_n_bytes = en_n_bytes
            new_en_n_len = en_n_len
            new_en_d_bytes = en_d_bytes
            new_en_d_len = en_d_len

        # Write out record:
        # Names (JP, EN, FR)
        new_data.extend(struct.pack("<h", jp_n_len))
        new_data.extend(jp_n_bytes)
        new_data.extend(struct.pack("<h", new_en_n_len))
        new_data.extend(new_en_n_bytes)
        new_data.extend(struct.pack("<h", fr_n_len))
        new_data.extend(fr_n_bytes)

        # Descs (JP, EN, FR)
        new_data.extend(struct.pack("<h", jp_d_len))
        new_data.extend(jp_d_bytes)
        new_data.extend(struct.pack("<h", new_en_d_len))
        new_data.extend(new_en_d_bytes)
        new_data.extend(struct.pack("<h", fr_d_len))
        new_data.extend(fr_d_bytes)

    # Save decrypted
    with open(bin_path + ".dec", "wb") as f:
        f.write(new_data)

    # Encrypt
    pad_len = 16 - (len(new_data) % 16)
    if pad_len < 16:
        new_data.extend(b"\x00" * pad_len)

    encrypted = encrypt_data(bytes(new_data))
    with open(bin_path, "wb") as f:
        f.write(encrypted)

    # Also update in patch folder
    patch_dest = os.path.join(r"D:\Viet Hoa Game\Setsuna_VietHoa_Patch\SETSUNA_Data\StreamingAssets\data\parameter", binary_file_name)
    os.makedirs(os.path.dirname(patch_dest), exist_ok=True)
    with open(patch_dest, "wb") as f:
        f.write(encrypted)

    print(f"Successfully packed and encrypted {binary_file_name} ({len(encrypted)} bytes)!")

def main():
    excel_path = r"D:\Viet Hoa Game\I_am_Setsuna-parameter-Steam.xlsx"
    print("Loading Excel...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    skill_files = [
        ("PlayerSkillDataMessage", "PlayerSkillDataMessage"),
        ("SetsunaSkillDataMessage", "SetsunaSkillDataMessage"),
        ("TsukushiSkillDataMessage", "TsukushiSkillDataMessage"),
        ("YomiSkillDataMessage", "YomiSkillDataMessage"),
        ("KishilSkillDataMessage", "KishilSkillDataMessage"),
        ("SionSkillDataMessage", "SionSkillDataMessage"),
        ("GrimreaperSkillDataMessage", "GrimreaperSkillDataMessage"),
        ("TwoPlayerCoopSkillDataMessage", "TwoPlayerCoopSkillDataMessage"),
        ("ThreePlayerCoopSkillDataMessag", "ThreePlayerCoopSkillDataMessage"),
        ("EnemySkillDataMessage", "EnemySkillDataMessage"),
        ("EnemyCoopSkillDataMessage", "EnemyCoopSkillDataMessage"),
        ("ItemSkillDataMessage", "ItemSkillDataMessage"),
    ]

    for s_name, f_name in skill_files:
        pack_skill_file(s_name, f_name, wb)

    print("\nALL SKILL FILES PACKED AND ENCRYPTED SUCCESSFULLY!")

if __name__ == "__main__":
    main()
