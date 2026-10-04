import openpyxl
import re

excel_path = "D:/Viet Hoa Game/I_am_Setsuna-parameter-Steam.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)

all_issues = {}

for name in wb.sheetnames:
    if name == "TOC":
        continue
    ws = wb[name]
    sheet_issues = {
        "hard_3plus_lines": [],
        "long_lines_32plus": [],
        "byte_overflow_234": [],
        "total_trans": 0,
    }

    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if not row or r_idx == 1:
            continue
        id_val = row[0]
        en_val = str(row[1]) if len(row) > 1 and row[1] is not None else ""
        vn_val = str(row[2]) if len(row) > 2 and row[2] is not None else ""

        if not vn_val.strip():
            continue

        sheet_issues["total_trans"] += 1
        pages = re.split(r"\n?\|\n?", vn_val)

        for p_i, page in enumerate(pages, start=1):
            p_bytes = page.encode("utf-16le")
            if len(p_bytes) > 234:
                sheet_issues["byte_overflow_234"].append({
                    "row": r_idx, "id": id_val, "page": p_i, "bytes": len(p_bytes)
                })

            lines = page.split("\n")
            if len(lines) > 2:
                sheet_issues["hard_3plus_lines"].append({
                    "row": r_idx, "id": id_val, "page": p_i, "num_lines": len(lines), "lines": lines
                })

            for l_i, line in enumerate(lines, start=1):
                vis_len = len(line)
                if vis_len > 32:
                    sheet_issues["long_lines_32plus"].append({
                        "row": r_idx, "id": id_val, "page": p_i, "line_idx": l_i, "len": vis_len, "line": line
                    })

    if sheet_issues["total_trans"] > 0:
        all_issues[name] = sheet_issues

header = "%-32s | %-10s | %-10s | %-16s | %-12s" % ("Sheet", "Da dich", "3+ Dong", "Dong > 32 ky tu", "> 234 Bytes")
print(header)
print("-" * len(header))

for name, s in all_issues.items():
    print("%-32s | %10d | %10d | %16d | %12d" % (
        name,
        s["total_trans"],
        len(s["hard_3plus_lines"]),
        len(s["long_lines_32plus"]),
        len(s["byte_overflow_234"]),
    ))

g_trans = sum(s["total_trans"] for s in all_issues.values())
g_hard = sum(len(s["hard_3plus_lines"]) for s in all_issues.values())
g_long = sum(len(s["long_lines_32plus"]) for s in all_issues.values())
g_byte = sum(len(s["byte_overflow_234"]) for s in all_issues.values())

print("-" * len(header))
print("%-32s | %10d | %10d | %16d | %12d" % (
    "TONG TOAN BO", g_trans, g_hard, g_long, g_byte
))
