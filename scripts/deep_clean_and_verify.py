import openpyxl
import re

def rewrap_to_pages(text, max_len=28):
    if not text:
        return ""
    # Flatten text first (remove existing \n and |)
    flat = text.replace("|", " ").replace("\r\n", " ").replace("\n", " ")
    flat = re.sub(r"\s+", " ", flat).strip()
    
    words = flat.split(" ")
    lines = []
    cur_line = []
    
    for w in words:
        # Test line with word
        test = " ".join(cur_line + [w])
        # Calculate visual length excluding tags like <NAME=...>
        clean_len = len(re.sub(r"<[^>]+>", "", test))
        if clean_len <= max_len:
            cur_line.append(w)
        else:
            if cur_line:
                lines.append(" ".join(cur_line))
            cur_line = [w]
            
    if cur_line:
        lines.append(" ".join(cur_line))
        
    # Group into pages of 2 lines max
    pages = []
    for i in range(0, len(lines), 2):
        page_lines = lines[i:i+2]
        pages.append("\n".join(page_lines))
        
    return "|".join(pages)

def main():
    excel_path = r"D:\i-am-setsuna\I_am_Setsuna-parameter-Steam.xlsx"
    wb = openpyxl.load_workbook(excel_path)
    
    # 1. Fix NormalConv Row 24
    ws_norm = wb["ScenarioMessageData_NormalConv"]
    row24_fixed = (
        "Chúng tôi sắp tổ chức lễ tiễn biệt\ncho người hiến tế.\n"
        "|\n"
        "Chính con gái tôi là người\nsẽ gánh vác trọng trách này...\n"
        "|\n"
        "Có lẽ đây cũng là định mệnh..."
    )
    ws_norm.cell(24, 3, row24_fixed)
    print("Fixed NormalConv Row 24!")
    
    # 2. Standardize line length across all scenario sheets
    scenario_sheets = [
        "ScenarioMessageData_Chapter_1",
        "ScenarioMessageData_Chapter_2",
        "ScenarioMessageData_Chapter_3",
        "ScenarioMessageData_Chapter_4",
        "ScenarioMessageData_NormalConv",
        "SubQuestMessageData"
    ]
    
    total_reformatted = 0
    for sname in scenario_sheets:
        ws = wb[sname]
        for r in range(2, ws.max_row + 1):
            vn = str(ws.cell(r, 3).value or "").strip()
            if not vn: continue
            
            # Re-wrap if any line exceeds 30 chars
            needs_wrap = False
            for page in vn.split("|"):
                for line in page.split("\n"):
                    clean = re.sub(r"<[^>]+>", "", line)
                    if len(clean) > 30:
                        needs_wrap = True
                        break
                if needs_wrap: break
                
            if needs_wrap:
                new_vn = rewrap_to_pages(vn, max_len=28)
                ws.cell(r, 3, new_vn)
                total_reformatted += 1
                
        print(f"Reformatted {sname}!")
        
    print(f"Total rows optimized for perfect UI display: {total_reformatted}")
    wb.save(excel_path)
    wb.save(r"D:\Viet Hoa Game\I_am_Setsuna-parameter-Steam.xlsx")
    print("Saved Master Excel in both locations successfully!")

if __name__ == "__main__":
    main()
