import openpyxl

excel_path = "D:/Viet Hoa Game/I_am_Setsuna-parameter-Steam.xlsx"
wb = openpyxl.load_workbook(excel_path)

# =========================================================================
# 1. PlayerSkillDataMessage (Endir)
# =========================================================================
ws = wb["PlayerSkillDataMessage"]
endir_skills = {
    6: "Tấn Công",
    8: "Cuồng Phong",
    9: "Đòn tấn công phạm vi rộng, có thể đẩy lùi kẻ địch.\nGây sát thương vật lý [Vô hệ] lên mọi kẻ địch quanh bản thân.",
    10: "Hào Quang",
    11: "Năng lượng trị thương hồi phục HP cho toàn bộ đồng minh\ngần mục tiêu. Giải trừ trạng thái bất lợi khi dùng ở chế độ Xung Lực.",
    12: "Chấn Động",
    13: "Giải phóng chùm ma lực tập trung, gây sát thương ma thuật\n[Vô hệ] lên toàn bộ kẻ địch gần mục tiêu.",
    14: "Bức Tường",
    15: "Dựng bức tường ma thuật quanh một mục tiêu, chống đỡ\ncả đòn tấn công vật lý lẫn ma thuật.",
    16: "Kiếm Khí",
    17: "Phóng ra lưỡi kiếm năng lượng gây sát thương vật lý\n[Vô hệ] lên toàn bộ kẻ địch trên đường thẳng.",
    18: "Tập Kích",
    19: "Lướt ra sau lưng kẻ địch trong chớp mắt, đánh úp bất ngờ.\nGây sát thương vật lý [Vô hệ] lên toàn bộ kẻ địch gần mục tiêu.",
    20: "Quang Mang",
    21: "Kết hợp năng lượng vật lý và ma thuật tạo vụ nổ gây\nsát thương đặc biệt [Vô hệ/4 đòn] lên kẻ địch gần mục tiêu.",
    22: "Ma Nguyên",
    23: "Niệm Tập Trung lên toàn bộ đồng minh gần mục tiêu.\nKhi ở trạng thái này, MP sẽ tự động hồi phục.",
    24: "Hồi Sinh",
    25: "Sức mạnh cứu rỗi hồi sinh đồng minh gần mục tiêu.\nĐồng thời hồi phục HP kể cả khi đồng minh chưa bị hạ gục.",
    26: "Trảm Kích",
    27: "Giải phóng toàn bộ sức mạnh vào đúng khoảnh khắc va chạm,\ngây sát thương vật lý [Vô hệ] lên toàn bộ kẻ địch gần mục tiêu.",
    28: "Giác Ngộ",
    29: "Lập tức làm đầy thanh SP của toàn bộ đồng minh gần mục tiêu,\nvà tăng tốc độ tích lũy SP của họ.",
    30: "Tia Lửa",
    31: "Gây sát thương vật lý [Vô hệ] lên kẻ địch gần mục tiêu.\nSát thương càng cao khi SP càng nhiều, và tích lũy lượng lớn SP.",
    32: "Chiến Hống",
    33: "Tăng mạnh Tấn Công, Phòng Thủ và tốc độ ATB của toàn đội.\nHiệu ứng chỉ duy trì 1 lượt nhưng tăng chỉ số cực lớn.",
    34: "Thiên Thạch",
    35: "Triệu hồi mưa thiên thạch gây sát thương ma thuật\n[Vô hệ] lên toàn bộ kẻ địch, bỏ qua 50% Phòng Thủ.",
    36: "Tăng Tốc",
    37: "Niệm Tăng ATB lên một mục tiêu. Sau khi hành động,\nthanh ATB sẽ được lấp đầy một phần ngay lập tức.",
    38: "Khởi Nguyên",
    39: "Gây sát thương đặc biệt [Đa hệ/8 đòn] lên kẻ địch gần mục tiêu.\nSát thương tăng theo số lần dùng Xung Lực và số lượng Flux tích lũy.",
}
for r, text in endir_skills.items():
    ws.cell(r, 3).value = text

# =========================================================================
# 2. SetsunaSkillDataMessage (Setsuna)
# =========================================================================
ws = wb["SetsunaSkillDataMessage"]
setsuna_skills = {
    6: "Tấn Công",
    8: "Trị Liệu",
    9: "Được kích hoạt bởi trái tim nhân từ, tăng cường ma lực\ntrị thương, hồi phục HP cho toàn bộ đồng minh gần mục tiêu.",
    10: "Lôi Điệp",
    11: "Dùng ma lực tạo ra tia sét, gây sát thương ma thuật\n[Quang] lên toàn bộ kẻ địch gần mục tiêu.",
    12: "Khiêu Khích",
    13: "Niệm Khiêu Khích lên một mục tiêu, khiến kẻ địch dồn\nmọi đòn tấn công vào nhân vật đó.",
    14: "Ban Lôi",
    15: "Thêm nguyên tố Quang vào đòn đánh của đồng minh gần mục tiêu.\nCác đòn tấn công hệ Quang cũng sẽ gây nhiều sát thương hơn.",
    16: "Thánh Quang",
    17: "Ném luân đao tích tụ ma lực, gây sát thương vật lý\n[Quang] lên toàn bộ kẻ địch trên một đường thẳng.",
    18: "Thanh Tẩy",
    19: "Giải trừ toàn bộ trạng thái bất lợi cho đồng minh gần mục tiêu.\nĐồng thời miễn nhiễm trạng thái bất lợi trong 2 lượt tiếp theo.",
    20: "Cổ Vũ",
    21: "Niệm Giải Phóng Liên Chiêu lên toàn đội, cho phép dùng\ncombo ngay cả khi chỉ có một nhân vật đầy thanh ATB.",
    22: "Cầu Nguyện",
    23: "Mài sắc tâm linh, hòa nhập cùng đất mẹ và hồi phục MP\ncho toàn bộ đồng minh gần mục tiêu.",
    24: "Giả Chết",
    25: "Người dùng giả chết khiến kẻ địch ngừng tấn công.\nĐồng thời tăng độ chính xác và tỉ lệ chí mạng cho lượt sau.",
    26: "Khích Lệ",
    27: "Hưởng ứng ý chí kiên định, tăng dũng khí và tinh thần chiến đấu,\ngia tăng Tấn Công cho toàn bộ đồng minh.",
    28: "Luân Vũ Trảm",
    29: "Dùng ma lực phóng luân đao xoay với tốc độ cực cao.\nGây sát thương vật lý [Vô hệ/6 đòn] lên kẻ địch gần mục tiêu.",
    30: "Trị Liệu II",
    31: "Được kích hoạt bởi tâm hồn sùng đạo, khuếch đại ma lực\ntrị thương của người dùng, hồi phục HP cho toàn bộ đồng minh.",
    32: "Hồi Sinh II",
    33: "Sức mạnh thần thánh hồi sinh toàn bộ đồng minh gần mục tiêu.\nĐồng thời hồi đầy HP kể cả khi đồng minh chưa bị hạ gục.",
    34: "Lôi Điệp II",
    35: "Triệu hồi vô số tia sét giáng xuống từ bầu trời,\ngây sát thương ma thuật [Quang] lên toàn bộ kẻ địch.",
    36: "Thánh Lực",
    37: "Gây sát thương ma thuật [Quang] lên kẻ địch gần mục tiêu.\nSát thương càng cao khi tổng lượng máu đã hồi phục càng lớn.",
    38: "Tỏa Sáng",
    39: "Hút toàn bộ ma lực xung quanh, gây sát thương ma thuật\n[Quang] lên toàn bộ kẻ địch và hồi phục HP cho toàn đội.",
}
for r, text in setsuna_skills.items():
    ws.cell(r, 3).value = text

# =========================================================================
# 3. SionSkillDataMessage (Aeterna)
# =========================================================================
ws = wb["SionSkillDataMessage"]
aeterna_skills = {
    6: "Tấn Công",
    8: "Phi Thân",
    9: "Nhảy vút lên không trung, tạm thời rời khỏi chiến trường.\nKhi đáp xuống, gây sát thương vật lý [Vô hệ] lên kẻ địch gần mục tiêu.",
    10: "Tiếp Ứng",
    11: "Vào thế thủ, khi đồng minh thực hiện đòn tấn công vật lý,\nsẽ lập tức bồi thêm đòn đánh gây sát thương vật lý [Vô hệ].",
    12: "Xuyên Tâm",
    13: "Gây sát thương đặc biệt [Vô hệ/15% Chí Mạng] lên kẻ địch\ngần mục tiêu. Đòn chí mạng sẽ gây thêm nhiều sát thương.",
    14: "Băng Kích",
    15: "Hạ nhiệt độ xung quanh xuống điểm đóng băng, gây sát thương\nma thuật [Thủy] lên toàn bộ kẻ địch gần mục tiêu.",
    16: "Ban Băng",
    17: "Thêm nguyên tố Thủy vào đòn đánh của đồng minh gần mục tiêu.\nCác đòn tấn công hệ Thủy cũng sẽ gây thêm nhiều sát thương.",
    18: "Băng Thương",
    19: "Gây sát thương vật lý [Thủy] lên toàn bộ kẻ địch trên đường\nthẳng. Có tỉ lệ gây trạng thái Đóng Băng.",
    20: "Cuồng Nộ",
    21: "Tăng mạnh Tấn Công của bản thân, nhưng đồng thời làm giảm\nPhòng Thủ của bản thân.",
    22: "Luân Xa",
    23: "Hồi phục HP và MP cho toàn bộ đồng minh gần bản thân.\nSức Mạnh người dùng càng cao, lượng HP/MP hồi phục càng nhiều.",
    24: "Băng Kích II",
    25: "Giải phóng ma lực ở nhiệt độ độ không tuyệt đối, gây sát thương\nma thuật [Thủy] lên toàn bộ kẻ địch.",
    26: "Âm Vang",
    27: "Vào thế thủ, khi đồng minh thực hiện đòn ma thuật, sẽ lập tức\nbồi thêm đòn đánh gây sát thương vật lý [Vô hệ].",
    28: "Toái Thương Kích",
    29: "Giải phóng xoáy năng lượng từ mũi thương, gây sát thương\nvật lý [Vô hệ/4 đòn] lên toàn bộ kẻ địch gần mục tiêu.",
    30: "Siêu Việt",
    31: "Hy sinh HP để đổi lấy hiệu ứng Bất Hạn và Giải Phóng Liên Chiêu,\nđồng thời làm đầy thanh SP ngay lập tức.",
    32: "Băng Tinh Trảm",
    33: "Đột kích từ trên cao kết hợp năng lượng vật lý và ma thuật,\ngây sát thương đặc biệt [Vô hệ] lên kẻ địch gần mục tiêu.",
    34: "Hoàng Gia Thương",
    35: "Tấn công bất ngờ, gây sát thương vật lý [Vô hệ/4 đòn]\nbỏ qua 50% Phòng Thủ lên toàn bộ kẻ địch gần mục tiêu.",
    36: "Tĩnh Tâm",
    37: "Dùng ma lực gạt bỏ mọi tạp niệm, gia tăng tỉ lệ chí mạng\ncho toàn bộ đồng minh gần mục tiêu.",
    38: "Nhật Thực",
    39: "Gây sát thương đặc biệt [Vô hệ] lên kẻ địch gần mục tiêu.\nSát thương càng tăng cao khi số lần đánh chí mạng càng nhiều.",
}
for r, text in aeterna_skills.items():
    ws.cell(r, 3).value = text

# =========================================================================
# 4. YomiSkillDataMessage (Nidr)
# =========================================================================
ws = wb["YomiSkillDataMessage"]
nidr_skills = {
    6: "Tấn Công",
    8: "Khiêu Khích",
    9: "Khiêu khích kẻ địch dồn đòn đánh vào bản thân. Đồng thời\ntăng Phòng Thủ cho bản thân khi dùng ở chế độ Xung Lực.",
    10: "Không Kích",
    11: "Nhảy vọt lên không trung rồi giáng xuống một đòn sấm sét,\ngây sát thương vật lý [Vô hệ] lên kẻ địch gần mục tiêu.",
    12: "Ngưng Tụ",
    13: "Truyền ma lực ngưng tụ vào vật phẩm đã chọn, gia tăng mạnh\nhiệu quả của vật phẩm đó.",
    14: "Sống Đao Trảm",
    15: "Đánh bằng sống đao, gây sát thương vật lý [Vô hệ] lên kẻ địch\ngần mục tiêu. Có tỉ lệ gây Tê Liệt.",
    16: "Phong Trảm",
    17: "Phóng sóng chân không từ kiếm áp, gây sát thương vật lý\n[Vô hệ] lên toàn bộ kẻ địch trên một đường thẳng.",
    18: "Phản Đòn",
    19: "Vào thế thủ, khi nhận đòn tấn công vật lý từ kẻ địch sẽ triệt tiêu\nsát thương và lập tức phản công.",
    20: "Hoàn Trả",
    21: "Vào thế thủ, khi nhận đòn ma thuật từ kẻ địch sẽ hoàn trả lại\nchiêu thức với lượng sát thương tăng cường.",
    22: "Thạch Trảm",
    23: "Gây sát thương vật lý [Vô hệ] lên toàn bộ kẻ địch gần mục tiêu.\nCó tỉ lệ gây Hóa Đá.",
    24: "Khuếch Tán",
    25: "Truyền ma lực khuếch tán vào vật phẩm đã chọn, giúp vật phẩm\náp dụng hiệu ứng lên toàn đội thay vì một người.",
    26: "Điểm Yếu Kích",
    27: "Gây sát thương vật lý [Vô hệ] lên toàn bộ kẻ địch gần mục tiêu.\nGây thêm nhiều sát thương khi đánh trúng điểm yếu.",
    28: "Kiên Cường",
    29: "Ban sức mạnh chịu đựng mọi nỗi đau, ngăn người dùng bị hạ gục\nkhi lượng HP tụt về 0.",
    30: "Cuồng Phong Đao",
    31: "Tạo lốc xoáy hút toàn bộ kẻ địch xung quanh lại gần.\nGây sát thương vật lý [Vô hệ] lên mọi kẻ địch gần bản thân.",
    32: "Ma Kích",
    33: "Gây sát thương vật lý [Vô hệ] lên kẻ địch gần mục tiêu.\nCó thể biến đổi nguyên tố trong đòn đánh của kẻ địch.",
    34: "Bộc Phát",
    35: "Gây sát thương vật lý [Vô hệ] lên kẻ địch gần mục tiêu.\nCàng có nhiều hiệu ứng tăng chỉ số, sát thương càng tăng cao.",
    36: "Phòng Thủ Bất Bại",
    37: "Tăng Phòng Thủ của bản thân, cho phép kích hoạt chế độ Xung Lực\nkhi bị tấn công để giảm thiểu toàn bộ sát thương.",
    38: "Phản Kháng",
    39: "Gây sát thương đặc biệt [Vô hệ] lên kẻ địch gần mục tiêu.\nCàng nhận nhiều lần sát thương, uy lực chiêu thức càng mạnh.",
}
for r, text in nidr_skills.items():
    ws.cell(r, 3).value = text

# =========================================================================
# 5. KishilSkillDataMessage (Kir)
# =========================================================================
ws = wb["KishilSkillDataMessage"]
kir_skills = {
    6: "Tấn Công",
    8: "Hỏa Cầu",
    9: "Ngưng tụ ma lực thành ngọn lửa, gây sát thương ma thuật\n[Hỏa] lên toàn bộ kẻ địch gần mục tiêu.",
    10: "Hút Máu",
    11: "Gây sát thương ma thuật [Vô hệ] lên kẻ địch gần mục tiêu,\nvà hấp thụ 100% sát thương gây ra thành HP.",
    12: "Ban Hỏa",
    13: "Thêm nguyên tố Hỏa vào đòn đánh của đồng minh gần mục tiêu.\nCác đòn tấn công hệ Hỏa cũng sẽ gây thêm nhiều sát thương.",
    14: "Hóa Giải",
    15: "Gây sát thương ma thuật [Vô hệ] lên toàn bộ kẻ địch gần mục tiêu.\nĐồng thời xóa bỏ toàn bộ hiệu ứng tăng chỉ số của chúng.",
    16: "Nham Thạch",
    17: "Dẫn truyền cơn thịnh nộ của đất mẹ, phóng dòng dung nham gây\nsát thương ma thuật [Hỏa/4 đòn] lên kẻ địch gần mục tiêu.",
    18: "Hồi Phục",
    19: "Niệm Hồi Phục lên toàn bộ đồng minh gần mục tiêu.\nKhi ở trạng thái này, HP sẽ tự động hồi phục liên tục.",
    20: "Độc Tố",
    21: "Gây sát thương ma thuật [Vô hệ] lên toàn bộ kẻ địch gần mục tiêu,\nđồng thời làm giảm chỉ số Tấn Công của chúng.",
    22: "Hút Năng Lượng",
    23: "Gây sát thương ma thuật [Vô hệ] lên kẻ địch gần mục tiêu,\nvà hấp thụ 100% sát thương gây ra thành MP.",
    24: "Hỏa Cầu II",
    25: "Sử dụng lượng lớn ma lực tạo ra vụ nổ lớn, gây sát thương\nma thuật [Hỏa] lên toàn bộ kẻ địch.",
    26: "Tàng Hình",
    27: "Niệm Vô Hình lên toàn bộ đồng minh gần mục tiêu.\nKhi ở trạng thái tàng hình, kẻ địch sẽ không thể tấn công.",
    28: "Luyện Kim Bại Hoại",
    29: "Gây sát thương ma thuật [Vô hệ] lên toàn bộ kẻ địch gần mục tiêu,\nđồng thời làm giảm chỉ số Phòng Thủ của chúng.",
    30: "Hộ Giáp",
    31: "Gia tăng kháng nguyên tố cho một mục tiêu.\nCó thể lựa chọn nguyên tố cụ thể muốn tăng kháng.",
    32: "Bộc Phá",
    33: "Tạo vụ nổ rực cháy gây sát thương ma thuật [Hỏa] lên kẻ địch\ngần mục tiêu, bỏ qua 50% chỉ số Phòng Thủ.",
    34: "Kích Hoạt Ma Lực",
    35: "Niệm Bất Hạn lên một mục tiêu, tăng lượng MP tiêu hao\nnhưng cường hóa mạnh mẽ uy lực của các kỹ năng.",
    36: "Suy Yếu",
    37: "Làm giảm kháng nguyên tố của một mục tiêu.\nCó thể lựa chọn nguyên tố cụ thể muốn giảm kháng.",
    38: "Diệt Vong",
    39: "Gây sát thương đặc biệt [Vô hệ] lên toàn bộ kẻ địch và đồng minh.\nCàng có nhiều đồng minh bị hạ gục, sát thương gây ra càng khủng khiếp.",
}
for r, text in kir_skills.items():
    ws.cell(r, 3).value = text

# =========================================================================
# 6. TsukushiSkillDataMessage (Julienne)
# =========================================================================
ws = wb["TsukushiSkillDataMessage"]
julienne_skills = {
    6: "Tấn Công",
    8: "Xung Kích",
    9: "Lao thẳng vào kẻ địch với siêu tốc độ cường hóa bởi ma lực,\ngây sát thương vật lý [Vô hệ] lên kẻ địch gần mục tiêu.",
    10: "Bảo Vệ",
    11: "Dùng ma lực tăng cường sức mạnh và khả năng miễn dịch,\ngia tăng Phòng Thủ cho toàn bộ đồng minh gần mục tiêu.",
    12: "Làm Chậm",
    13: "Giảm tốc độ của toàn bộ kẻ địch gần mục tiêu.\nKhiến thanh ATB của chúng nạp chậm hơn, giảm tần suất hành động.",
    14: "Trọng Lực",
    15: "Tạo khối cầu trọng lực hút kẻ địch về trung tâm, gây sát thương\nma thuật [Thời Gian] lên toàn bộ kẻ địch gần mục tiêu.",
    16: "Ban Thời",
    17: "Thêm nguyên tố Thời Gian vào đòn đánh của đồng minh gần mục tiêu.\nCác đòn đánh hệ Thời Gian cũng sẽ gây thêm nhiều sát thương.",
    18: "Hoàn Bích Kích",
    19: "Gây sát thương vật lý [Vô hệ] lên toàn bộ kẻ địch gần mục tiêu.\nSau khi dùng, thanh ATB sẽ được làm đầy một phần ngay lập tức.",
    20: "Tăng Tốc",
    21: "Gia tăng tốc độ cho toàn bộ đồng minh gần mục tiêu.\nKhiến thanh ATB nạp nhanh hơn, tăng tần suất hành động.",
    22: "Phá Mộng",
    23: "Gây sát thương vật lý [Vô hệ] lên toàn bộ kẻ địch gần mục tiêu.\nCó tỉ lệ gây Hỗn Loạn và Choáng.",
    24: "Ngưng Đọng",
    25: "Can thiệp vào sự cân bằng của đất trời, thay đổi dòng thời gian\nvà gây trạng thái Ngưng Đọng lên toàn bộ kẻ địch gần mục tiêu.",
    26: "Hư Vô",
    27: "Kích hoạt sức mạnh bẻ cong không gian, gây sát thương vật lý\n[Vô hệ] bỏ qua 50% Phòng Thủ lên toàn bộ kẻ địch.",
    28: "Vạn Vật",
    29: "Thao túng thực tại, cưỡng ép kích hoạt Điểm Dị Thường và gây\nsát thương đặc biệt [Vô hệ] lên toàn bộ kẻ địch gần mục tiêu.",
    30: "Vĩnh Cửu",
    31: "Làm chậm dòng chảy ma lực, kéo dài thời gian hiệu lực của các\nhiệu ứng tăng chỉ số cho toàn bộ đồng minh gần mục tiêu.",
    32: "Cuồng Loạn",
    33: "Phá vỡ dòng chảy thời gian, tăng tốc độ xuất chiêu cực đại.\nGây sát thương vật lý [Vô hệ/8 đòn] lên kẻ địch gần mục tiêu.",
    34: "Ảo Ảnh",
    35: "Làm rối loạn kẻ địch bằng ảo ảnh tạo từ ma lực,\nban khả năng Né Đòn Tuyệt Đối cho một mục tiêu.",
    36: "Oa Phách",
    37: "Gây sát thương đặc biệt [Vô hệ] lên toàn bộ kẻ địch.\nLượng HP hiện tại càng thấp, sát thương gây ra càng lớn.",
    38: "Luân Hồi Kỷ Điểm",
    39: "Gây sát thương đặc biệt [Vô hệ/8 đòn] lên kẻ địch gần mục tiêu.\nSát thương tăng theo tổng số hành động và số loại hành động đã thực hiện.",
}
for r, text in julienne_skills.items():
    ws.cell(r, 3).value = text

# =========================================================================
# 7. GrimreaperSkillDataMessage (Fides)
# =========================================================================
ws = wb["GrimreaperSkillDataMessage"]
fides_skills = {
    6: "Tấn Công",
    8: "Xoáy Tử Thần",
    9: "Tận dụng tối đa tầm quét của lưỡi hái, gây sát thương\nđặc biệt [Vô hệ/2 đòn] lên toàn bộ kẻ địch gần bản thân.",
    10: "Hắc Vụ",
    11: "Dùng ma lực khuếch đại sức mạnh bóng tối, gây sát thương\nma thuật [Hắc Ám] lên toàn bộ kẻ địch gần mục tiêu.",
    12: "Nguyền Trảm",
    13: "Gây sát thương vật lý [Hắc Ám] lên toàn bộ kẻ địch gần bản thân.\nCó tỉ lệ gây trạng thái Suy Kiệt.",
    14: "Ban Hắc",
    15: "Thêm nguyên tố Hắc Ám vào đòn đánh của đồng minh gần mục tiêu.\nCác đòn đánh hệ Hắc Ám cũng sẽ gây thêm nhiều sát thương.",
    16: "Phá Địa Trảm",
    17: "Năng lượng tuôn trào từ lưỡi hái xé toạc mặt đất, gây sát thương\nvật lý [Vô hệ] lên toàn bộ kẻ địch gần mục tiêu.",
    18: "Ngục Hỏa",
    19: "Tập kích từ phía sau, gây sát thương vật lý [Hắc Ám] lên kẻ địch\ngần mục tiêu. Có tỉ lệ gây Tử Vong tức thì.",
    20: "Hấp Thu",
    21: "Niệm Hấp Thu HP và Hấp Thu MP lên toàn đội gần mục tiêu,\ngiúp họ hút HP/MP khi thực hiện tấn công.",
    22: "Đại Dịch",
    23: "Gây sát thương vật lý [Hắc Ám] lên toàn bộ kẻ địch gần mục tiêu.\nCàng có nhiều trạng thái bất lợi, sát thương càng tăng cao.",
    24: "Ảnh Tung",
    25: "Niệm Ảo Ảnh lên một mục tiêu. Khi ở trạng thái này, nhân vật\nsẽ tự động bồi thêm đòn đánh phụ khi dùng lệnh 'Tấn Công'.",
    26: "Đoạt Hồn",
    27: "Giảm Tấn Công và Phòng Thủ của toàn bộ kẻ địch gần mục tiêu,\nđồng thời gia tăng Tấn Công và Phòng Thủ cho bản thân.",
    28: "Ảo Ảnh Ma Quái",
    29: "Gây hàng loạt trạng thái bất lợi (Suy Kiệt, Tê Liệt, Hỗn Loạn\nvà Choáng) lên toàn bộ kẻ địch trên một đường thẳng.",
    30: "Hố Đen",
    31: "Tạo ra khối cầu bóng tối hút kẻ địch vào trong, gây sát thương\nma thuật [Hắc Ám] lên toàn bộ kẻ địch gần mục tiêu.",
    32: "Hắc Khế",
    33: "Hy sinh HP chuyển hóa thành sức mạnh hủy diệt. Gây sát thương\nvật lý [Hắc Ám] lên toàn bộ kẻ địch trên một đường thẳng.",
    34: "Hắc Thể",
    35: "Vật chất tối tạo từ ma lực gây sát thương ma thuật\n[Hắc Ám/6 đòn] lên toàn bộ kẻ địch.",
    36: "Long Hồn",
    37: "Chế độ Xung Lực bị vô hiệu hóa, nhưng đổi lại toàn bộ\nchỉ số của nhân vật được gia tăng vượt bậc.",
    38: "Tận Thế Rồng",
    39: "Gây sát thương đặc biệt [Vô hệ] lên kẻ địch gần mục tiêu.\nCàng có nhiều điểm Xung Lực chưa sử dụng, sát thương càng khủng khiếp.",
}
for r, text in fides_skills.items():
    ws.cell(r, 3).value = text

wb.save(excel_path)
print("Updated all 7 Player/Ally Skill sheets successfully (231 entries)!")
