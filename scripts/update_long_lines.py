import openpyxl
import json
import re

UPDATES = {
    307: "Nếu anh ta giỏi đến mức khiến người ta dè chừng,\nthì chắc chắn sẽ là một người bạn đồng hành\nvô cùng đáng tin cậy...",
    1020: "Dù sao quân của ta cũng sẽ sớm tìm ra anh ta thôi.\nNgay khi thấy, ta sẽ lệnh cho phi thuyền\nsẵn sàng xuất phát ngay lập tức!",
    1111: "Chuyến hành trình của các cháu sẽ vô cùng hiểm trở,\nvà các cháu còn chẳng biết liệu nó\ncó được đền đáp hay không...",
    1138: "Thương hại tôi á? Tôi có thể yếu, nhưng tôi chẳng cần ai thương hại cả!\nTôi sống cuộc đời của TÔI theo cách của TÔI!",
    1180: "Chậc! Xem ra bị các người bắt thóp rồi!\nĐúng thế đấy...\nTa chính là <NAME=NPC_10080>, và <NAME=NPC_10080> chính là ta!",
    1187: "Vì cựu Lãnh chúa thì được... chứ lái phi thuyền cho tên hiện tại thì không bao giờ nhé! Không đời nào!",
    1242: "Nhưng ta không ngờ ngươi lại tự dâng tới tận miệng thế này!\nNgươi đã giúp ta đỡ tốn khối công sức đấy...",
    1304: "Không sao, chừng đó là quá đủ rồi.\nCái tên Lãnh chúa khốn kiếp đó... dám lạm quyền\nlàm những chuyện đồi bại thế này...",
    1326: "Không phải ai cũng làm con bài đàm phán được đâu...\nNhưng hôm nay, ta lại có một lựa chọn vô cùng tuyệt hảo!",
    1353: "Không còn thời gian đâu! Giờ ta chỉ có thể\ncố hết sức để giảm thiểu thiệt hại nhiều nhất có thể thôi..."
}

def main():
    excel_path = r"d:\Viet Hoa Game\I_am_Setsuna-parameter-Steam.xlsx"
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["ScenarioMessageData_Chapter_1"]
    
    updated_count = 0
    for r in range(3, 1406):
        id_val = ws.cell(r, 1).value
        m = re.search(r'"Id":(\d+)', str(id_val))
        if m:
            rid = int(m.group(1))
            if rid in UPDATES:
                new_text = UPDATES[rid]
                ws.cell(r, 3).value = new_text
                byte_len = len(new_text.encode('utf-16le'))
                print(f"Updated RID {rid} (Row {r}): {byte_len} bytes")
                updated_count += 1
                
    wb.save(excel_path)
    print(f"Saved {updated_count} updates to {excel_path}")
    
    # Also update JSON
    json_path = r"d:\Viet Hoa Game\translation_export\ScenarioMessageData_Chapter_1.json"
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
        
    j_count = 0
    for item in data:
        id_str = item.get("id", "")
        m = re.search(r'"Id":(\d+)', id_str)
        if m:
            rid = int(m.group(1))
            if rid in UPDATES:
                item["vietnamese"] = UPDATES[rid]
                j_count += 1
                
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"Saved {j_count} updates to {json_path}")

if __name__ == "__main__":
    main()
