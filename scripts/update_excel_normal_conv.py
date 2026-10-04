import openpyxl
import json

excel_path = "D:/Viet Hoa Game/I_am_Setsuna-parameter-Steam.xlsx"
wb = openpyxl.load_workbook(excel_path)
ws = wb["ScenarioMessageData_NormalConv"]

with open("D:/Viet Hoa Game/temp_scripts/normal_conv_translated.json", "r", encoding="utf-8") as f:
    trans = json.load(f)

with open("D:/Viet Hoa Game/temp_scripts/normal_conv_all.json", "r", encoding="utf-8") as f:
    all_records = json.load(f)

# Clear old rows and write header
ws.delete_rows(2, ws.max_row)

ws.cell(1, 1).value = "ID"
ws.cell(1, 2).value = "English"
ws.cell(1, 3).value = "Vietnamese"

for r_idx, rec in enumerate(all_records, start=2):
    rid = rec["rid"]
    npc = rec["npc"]
    quest = rec["quest"]
    id_str = f'{rid}_{{"NpcNm":"{npc}","QuestNm":"{quest}","Head":{{"Type":0,"Id":{rid},"NpcId":0,"QuestId":0}}}}'
    en_text = rec["en"]
    vn_text = trans.get(str(rid), en_text)
    
    ws.cell(r_idx, 1).value = id_str
    ws.cell(r_idx, 2).value = en_text
    ws.cell(r_idx, 3).value = vn_text

wb.save(excel_path)
print(f"Updated ScenarioMessageData_NormalConv in Excel with {len(all_records)} rows!")
