import openpyxl
import re

def clean_dialogue(text, max_chars=27):
    if not text:
        return ""
    # Normalize separators
    flat = text.replace("\r\n", " ").replace("\n", " ").replace("|", " ")
    flat = re.sub(r"\s+", " ", flat).strip()
    
    words = flat.split(" ")
    lines = []
    cur = []
    
    for w in words:
        cand = " ".join(cur + [w])
        clean = re.sub(r"<[^>]+>", "", cand)
        if len(clean) <= max_chars:
            cur.append(w)
        else:
            if cur:
                lines.append(" ".join(cur))
            cur = [w]
    if cur:
        lines.append(" ".join(cur))
        
    pages = []
    for i in range(0, len(lines), 2):
        p_lines = lines[i:i+2]
        pages.append("\n".join(p_lines))
        
    return "|".join(pages)

def main():
    excel_path = r"D:\i-am-setsuna\I_am_Setsuna-parameter-Steam.xlsx"
    wb = openpyxl.load_workbook(excel_path)
    
    scenario_sheets = [
        "ScenarioMessageData_Chapter_1",
        "ScenarioMessageData_Chapter_2",
        "ScenarioMessageData_Chapter_3",
        "ScenarioMessageData_Chapter_4",
        "ScenarioMessageData_NormalConv",
        "SubQuestMessageData"
    ]
    
    for sname in scenario_sheets:
        ws = wb[sname]
        for r in range(2, ws.max_row + 1):
            vn = str(ws.cell(r, 3).value or "").strip()
            if vn:
                ws.cell(r, 3, clean_dialogue(vn, max_chars=27))
        print(f"Perfected {sname}!")
        
    wb.save(excel_path)
    wb.save(r"D:\Viet Hoa Game\I_am_Setsuna-parameter-Steam.xlsx")
    print("Saved clean dialogue to both Excels!")

if __name__ == "__main__":
    main()
