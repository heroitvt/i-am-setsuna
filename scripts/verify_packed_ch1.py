import openpyxl
import re
import struct
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
from cryptography.hazmat.backends import default_backend

KEY = b"8xTD|EgD|b?07QDj"
IV = b"/]s@*CxLzM!9Qd%("

def decrypt_data(data):
    cipher = Cipher(algorithms.AES(KEY), modes.CBC(IV), backend=default_backend())
    decryptor = cipher.decryptor()
    return decryptor.update(data) + decryptor.finalize()

def parse_string_field(data):
    if len(data) == 0:
        return ""
    num_pages = data[0]
    pages = []
    cur = 1
    for _ in range(num_pages):
        plen = data[cur]; cur += 1
        p_bytes = data[cur:cur+plen]; cur += plen
        pages.append(p_bytes.decode('utf-16le'))
    if cur != len(data):
        raise ValueError(f"Consumed {cur} != {len(data)}")
    return "\n|\n".join(pages)

def main():
    packed_file = r"d:\Viet Hoa Game\temp_scripts\ScenarioMessageData_Chapter_1.packed"
    excel_path = r"d:\Viet Hoa Game\I_am_Setsuna-parameter-Steam.xlsx"

    print("Loading Excel for verification...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["ScenarioMessageData_Chapter_1"]
    excel_map = {}
    for r in range(3, 1406):
        id_val = ws.cell(r, 1).value
        vn_val = ws.cell(r, 3).value
        m = re.search(r'"Id":(\d+)', str(id_val))
        if m:
            rid = int(m.group(1))
            excel_map[rid] = str(vn_val)

    print("Decrypting and parsing packed binary...")
    with open(packed_file, "rb") as f:
        enc_data = f.read()
    dec = decrypt_data(enc_data)

    offset = 0
    records_count = 0
    errors = []

    for r in range(1403):
        if offset + 16 > len(dec):
            errors.append(f"Unexpected EOF at offset {offset} before record {r}")
            break
        magic, rid, l1, l2, l3, rlen, npcid = struct.unpack("<ihhhhHh", dec[offset:offset+16])
        if magic != -1:
            errors.append(f"Record {r} magic is {magic} != -1")
        if rid != r:
            errors.append(f"Record {r} rid is {rid} != {r}")
            
        cur = offset + 16
        n_len = dec[cur]; cur += 1
        npc = dec[cur:cur+n_len].decode('utf-16le', errors='ignore'); cur += n_len
        
        q_len = dec[cur]; cur += 1
        quest = dec[cur:cur+q_len].decode('utf-16le', errors='ignore'); cur += q_len
        
        s1 = dec[cur:cur+l1]; cur += l1
        s2 = dec[cur:cur+l2]; cur += l2
        s3 = dec[cur:cur+l3]; cur += l3
        
        calc_rlen = 16 + 1 + n_len + 1 + q_len + l1 + l2 + l3
        if calc_rlen != rlen:
            errors.append(f"Record {rid} rlen mismatch: calc {calc_rlen} != header {rlen}")
            
        # Parse s2
        decoded_vn = parse_string_field(s2)
        # Compare with excel
        expected_vn = excel_map.get(rid, "")
        # Normalize line endings and page breaks for comparison
        norm_decoded = re.sub(r'\r?\n?\|\r?\n?', '|', decoded_vn).replace('\r\n', '\n').strip()
        norm_expected = re.sub(r'\r?\n?\|\r?\n?', '|', expected_vn).replace('\r\n', '\n').strip()
        
        if norm_decoded != norm_expected:
            errors.append(f"Record {rid} text mismatch:\n  Decoded:  {norm_decoded[:50]}\n  Expected: {norm_expected[:50]}")
            
        # Check tags
        expected_tags = re.findall(r'<[^>]+>', expected_vn)
        decoded_tags = re.findall(r'<[^>]+>', decoded_vn)
        if expected_tags != decoded_tags:
            errors.append(f"Record {rid} tag mismatch: {decoded_tags} != {expected_tags}")
            
        records_count += 1
        offset += rlen

    pad_bytes = len(dec) - offset
    print(f"Verified records: {records_count}/1403")
    print(f"Padding at end: {pad_bytes} bytes")
    print(f"Total validation errors: {len(errors)}")
    if errors:
        for e in errors[:10]:
            print("  ERROR:", e)
    else:
        print("PERFECT! 100% of 1403 records verified with 0 errors!")

if __name__ == "__main__":
    main()
