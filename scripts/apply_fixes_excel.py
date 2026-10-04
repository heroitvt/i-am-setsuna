import openpyxl
import re

excel_path = "D:/Viet Hoa Game/I_am_Setsuna-parameter-Steam.xlsx"
wb = openpyxl.load_workbook(excel_path)

ch1_fixes = {
    43: 'Bạn nhận "<NAME=ITEM_LWS_101>"\ntừ <NAME=NPC_10020>.',
    49: 'Một quầng sáng...\nCó luồng ma lực tỏa ra từ nó.',
    60: 'Kẻ địch nào cũng có điểm yếu.\nNgươi phải nắm bắt thời cơ!',
    61: 'Đây là công việc cuối cùng,\n<NAME=CP_0001>, ta trông cậy ở ngươi!',
    103: 'Thật vui vì nhiệm vụ cuối cùng\nta được đi cùng một người như ngươi.',
    126: 'Ngoài khơi xa có một hòn đảo,\nvà trên đảo có một ngôi làng nhỏ...\n|\nNơi đó có một thiếu nữ tròn 18 tuổi\nvào năm nay.',
    129: 'Cứ 10 năm, một vật tế được chọn\nđến Vùng Đất Tận Cùng làm nhiệm vụ.',
    166: 'Ta sẽ bảo vệ <NAME=CP_0002>,\ndù có phải hy sinh tính mạng!',
    175: 'Hai pháp sư đang giữ chân ngươi đấy,\nngươi biết không?',
    186: 'Bao đời nay, vật tế của làng này\nđã hy sinh để bảo vệ thế giới...',
    193: 'Đúng là lính đánh thuê...\nChỉ quan tâm điều mình cần biết...',
    199: 'Thôi nào, giả bộ quan tâm chút đi!\nĐúng là đồ vô cảm...',
    200: 'Ta đùa thôi, làng không có máy chém đâu.\nNhưng ngươi may mắn đấy...',
    210: 'Chính vật tế đã xin tha cho cậu,\nnên chẳng ai phản đối được.',
    274: '<NAME=CP_0002> đã gia nhập đội ngũ.',
    276: 'Chưa quyết định xong chuyện của ngươi,\nbọn ta không thể thả ngươi đi.',
    287: 'Trong góc khuất lịch sử, vẫn luôn có\nnhững người lặng lẽ bảo vệ nhân loại...',
    294: 'Nghi lễ hiến tế lặp đi lặp lại\nlà để mang lại tương lai cho nhân loại...',
    296: 'Ta nghe nói cậu đã đánh đuổi\nlũ quái vật tấn công làng, đúng không?',
    310: 'Người giỏi đến mức nguy hiểm như anh ta\nchắc chắn sẽ là bạn đồng hành đáng tin...',
    398: '<NAME=CP_0002> nhận "<NAME=ITEM_EVT_001>"\ntừ <NAME=NPC_10010>.',
    491: 'Lạ thật... Làng bỏ hoang thế này\nđáng lẽ phải đầy rẫy quái vật chứ...',
    558: 'Chính là tôi đây! Xin mời lại đây!\nThưa quý ông quý bà và các bạn nhỏ!\n|\nHãy xem tôi xoay chiếc đĩa này đây...',
    1023: 'Quân của ta sẽ tìm ra anh ta ngay.\nPhi thuyền sẽ sẵn sàng cho các người!',
    1114: 'Hành trình sẽ vô cùng gian nan,\nvà chưa chắc sẽ được đền đáp...',
    1127: 'Nhận được "<NAME=ITEM_EVT_002>"\ntừ <NAME=NPC_10070>.',
    1183: 'Bị bắt thóp rồi! Đúng thế đấy...\nTa chính là <NAME=NPC_10080>!',
    1307: 'Thế là quá đủ rồi. Tên Lãnh chúa đó...\ndám lạm quyền làm trò đồi bại thế này...',
}

subquest_fixes = {
    3: 'Aha! Ngươi là <NAME=CP_0001>\ncủa tộc Mặt Nạ phải không!\n|\nTa có thư gửi cho ngươi đây!',
    9: 'Ở vùng này có ba bệ thờ trọng yếu.\n|\nNếu giơ viên spritnite trước bệ thờ,\nngươi sẽ được dẫn lối cho tương lai.',
}

ws_ch1 = wb["ScenarioMessageData_Chapter_1"]
for row_idx, new_text in ch1_fixes.items():
    ws_ch1.cell(row_idx, 3).value = new_text

ws_sq = wb["SubQuestMessageData"]
for row_idx, new_text in subquest_fixes.items():
    ws_sq.cell(row_idx, 3).value = new_text

wb.save(excel_path)
print("Updated all 30 rows in Master Excel successfully!")
