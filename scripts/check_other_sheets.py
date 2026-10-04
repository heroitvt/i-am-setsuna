import openpyxl

def main():
    wb = openpyxl.load_workbook(r"d:\Viet Hoa Game\I_am_Setsuna-parameter-Steam.xlsx", data_only=True)
    ws = wb["ScenarioMessageData_NormalConv"]
    print(f"ScenarioMessageData_NormalConv max_row: {ws.max_row}")
    
    translated = 0
    empty = 0
    for r in range(3, ws.max_row + 1):
        vn = ws.cell(r, 3).value
        if vn and str(vn).strip():
            translated += 1
        else:
            empty += 1
            
    print(f"Normal Conversation translated: {translated}, empty: {empty}")
    
    # Also check other sheets!
    for sname in ["TitleMessage", "SystemMessage", "BattleMessage", "CampMessage"]:
        if sname in wb.sheetnames:
            sws = wb[sname]
            tr = sum(1 for r in range(2, sws.max_row+1) if sws.cell(r, 3).value and str(sws.cell(r, 3).value).strip())
            em = sum(1 for r in range(2, sws.max_row+1) if not (sws.cell(r, 3).value and str(sws.cell(r, 3).value).strip()))
            print(f"{sname}: max_row={sws.max_row}, translated={tr}, empty={em}")
        else:
            print(f"{sname} NOT in Excel!")

if __name__ == "__main__":
    main()
