import openpyxl

def main():
    wb = openpyxl.load_workbook(r"d:\Viet Hoa Game\I_am_Setsuna-parameter-Steam.xlsx", data_only=True)
    ws = wb["ScenarioMessageData_Chapter_1"]
    for row in range(120, 130):
        rid = ws.cell(row=row, column=1).value
        s2 = ws.cell(row=row, column=5).value
        print(f"Row {row}: rid={rid}, s2={repr(str(s2)[:40])}")

if __name__ == "__main__":
    main()
