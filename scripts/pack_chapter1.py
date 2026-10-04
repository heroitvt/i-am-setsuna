import openpyxl
import re
import struct
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
from cryptography.hazmat.backends import default_backend

KEY = b"8xTD|EgD|b?07QDj"
IV = b"/]s@*CxLzM!9Qd%("

def decrypt_data(data):
    cipher = Cipher(algorithms.AES(KEY), modes.CBC(IV), backend=default_backend())
    decryptor = cipher.decryptor()
    return decryptor.update(data) + decryptor.finalize()

def encrypt_data(data):
    cipher = Cipher(algorithms.AES(KEY), modes.CBC(IV), backend=default_backend())
    encryptor = cipher.encryptor()
    return encryptor.update(data) + encryptor.finalize()

def encode_string_field(text):
    if text is None or str(text).strip() == "":
        return b""
    text_str = str(text)
    if "|" in text_str:
        pages = re.split(r'\n?\|\n?', text_str)
    else:
        pages = [text_str]
        
    num_pages = len(pages)
    buf = bytearray()
    buf.append(num_pages)
    for p in pages:
        p_bytes = p.encode('utf-16le')
        if len(p_bytes) > 255:
            raise ValueError(f"Page byte length {len(p_bytes)} > 255 for text: {repr(p[:30])}")
        buf.append(len(p_bytes))
        buf.extend(p_bytes)
    return bytes(buf)

def main():
    excel_path = r"d:\Viet Hoa Game\I_am_Setsuna-parameter-Steam.xlsx"
    pristine_path = r"d:\Viet Hoa Game\I am Setsuna\SETSUNA_Data\StreamingAssets\data\parameter_backup_before_MASTER_FINAL\ScenarioMessageData_Chapter_1.backup_before_CH1_LOOSE_TEST"
    output_path = r"d:\Viet Hoa Game\temp_scripts\ScenarioMessageData_Chapter_1.packed"

    print("Loading Excel translations...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["ScenarioMessageData_Chapter_1"]
    
    translations = {}
    for r in range(3, 1406):
        id_val = ws.cell(r, 1).value
        vn_val = ws.cell(r, 3).value
        m = re.search(r'"Id":(\d+)', str(id_val))
        if m:
            rid = int(m.group(1))
            translations[rid] = vn_val
    print(f"Loaded {len(translations)} translations.")

    print("Reading pristine parameter file...")
    with open(pristine_path, "rb") as f:
        encrypted_pristine = f.read()
    dec = decrypt_data(encrypted_pristine)
    print(f"Pristine decrypted size: {len(dec)} bytes.")

    new_decrypted = bytearray()
    offset = 0
    records_processed = 0

    for r in range(1403):
        magic, rid, l1, l2, l3, rlen, npcid = struct.unpack("<ihhhhHh", dec[offset:offset+16])
        cur = offset + 16
        
        npc_len = dec[cur]; cur += 1
        npc_bytes = dec[cur:cur+npc_len]; cur += npc_len
        
        q_len = dec[cur]; cur += 1
        q_bytes = dec[cur:cur+q_len]; cur += q_len
        
        s1_bytes = dec[cur:cur+l1]; cur += l1
        orig_s2_bytes = dec[cur:cur+l2]; cur += l2
        s3_bytes = dec[cur:cur+l3]; cur += l3
        
        # Build new s2 from Vietnamese translation
        vn_text = translations.get(rid)
        if vn_text is None:
            raise ValueError(f"Missing translation for RID {rid}")
            
        new_s2_bytes = encode_string_field(vn_text)
        new_l2 = len(new_s2_bytes)
        new_rec_len = 16 + 1 + npc_len + 1 + q_len + l1 + new_l2 + l3
        
        # Build record header
        rec_header = struct.pack("<ihhhhHh", magic, rid, l1, new_l2, l3, new_rec_len, npcid)
        
        # Append all parts
        new_decrypted.extend(rec_header)
        new_decrypted.append(npc_len)
        new_decrypted.extend(npc_bytes)
        new_decrypted.append(q_len)
        new_decrypted.extend(q_bytes)
        new_decrypted.extend(s1_bytes)
        new_decrypted.extend(new_s2_bytes)
        new_decrypted.extend(s3_bytes)
        
        records_processed += 1
        offset += rlen

    print(f"Processed {records_processed} records.")
    print(f"New unpadded size: {len(new_decrypted)} bytes.")
    
    # Pad to multiple of 16 bytes
    pad_len = (16 - (len(new_decrypted) % 16)) % 16
    if pad_len > 0:
        new_decrypted.extend(b"\x00" * pad_len)
    print(f"Padded size: {len(new_decrypted)} bytes (padding: {pad_len} bytes).")

    # Encrypt
    encrypted_new = encrypt_data(bytes(new_decrypted))
    print(f"Encrypted size: {len(encrypted_new)} bytes.")

    with open(output_path, "wb") as f:
        f.write(encrypted_new)
    print(f"Successfully wrote packed file to {output_path}")

if __name__ == "__main__":
    main()
