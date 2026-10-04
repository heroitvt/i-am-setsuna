import os
import struct
import openpyxl
import re
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
from cryptography.hazmat.backends import default_backend

KEY = b"8xTD|EgD|b?07QDj"
IV = b"/]s@*CxLzM!9Qd%("

def encrypt_data(data):
    cipher = Cipher(algorithms.AES(KEY), modes.CBC(IV), backend=default_backend())
    encryptor = cipher.encryptor()
    return encryptor.update(data) + encryptor.finalize()

def decrypt_data(data):
    cipher = Cipher(algorithms.AES(KEY), modes.CBC(IV), backend=default_backend())
    decryptor = cipher.decryptor()
    return decryptor.update(data) + decryptor.finalize()

def main():
    excel_path = r"d:\Viet Hoa Game\I_am_Setsuna-parameter-Steam.xlsx"
    cpk_path = r"d:\Viet Hoa Game\I am Setsuna\SETSUNA_Data\StreamingAssets\x86_64\parameter.cpk"

    print("1. Loading Excel translations for Chapter 1...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["ScenarioMessageData_Chapter_1"]

    translations = {}
    for r in range(3, ws.max_row + 1):
        id_val = ws.cell(r, 1).value
        vn_val = ws.cell(r, 3).value
        if id_val and vn_val:
            m = re.search(r'"Id":(\d+)', str(id_val))
            if m:
                rid = int(m.group(1))
                translations[rid] = str(vn_val).strip()

    print(f"Loaded {len(translations)} translations from Excel.")

    # Read the exact 372,704 bytes at offset 9605120 from CPK
    offset = 9605120
    file_size = 372704

    with open(cpk_path, "rb") as f:
        f.seek(offset)
        raw_enc = f.read(file_size)

    assert len(raw_enc) == file_size, f"Expected {file_size} bytes, got {len(raw_enc)}"

    dec = bytearray(decrypt_data(raw_enc))
    assert len(dec) == file_size

    pos = 0
    records_count = 0
    updated_count = 0
    too_long_count = 0

    while pos < len(dec):
        if pos + 16 > len(dec): break
        magic, rid, l1, l2, l3, rlen, npcid = struct.unpack("<ihhhhHh", dec[pos:pos+16])
        if rlen == 0 or pos + rlen > len(dec): break

        cur = pos + 16
        n_len = dec[cur]; cur += 1
        npc_name = dec[cur:cur+n_len].decode("ascii", errors="ignore"); cur += n_len
        q_len = dec[cur]; cur += 1
        quest_name = dec[cur:cur+q_len].decode("ascii", errors="ignore"); cur += q_len

        cur += l1 # skip s1 (Japanese)
        s2_start = cur
        cur += l2 # old s2 slot

        if rid in translations:
            vn_text = translations[rid]
            pages = vn_text.split("\n|\n") if "\n|\n" in vn_text else vn_text.split("|")
            
            # Encode pages
            new_s2_buf = bytearray()
            new_s2_buf.append(len(pages))
            for p in pages:
                p_clean = p.strip()
                p_u16 = p_clean.encode("utf-16le")
                new_s2_buf.append(len(p_u16))
                new_s2_buf.extend(p_u16)

            # Check if fits into old l2 slot
            old_slot_len = l2
            new_len = len(new_s2_buf)

            if new_len <= old_slot_len:
                # Pad with spaces in UTF-16LE or nulls inside the text / slot
                diff = old_slot_len - new_len
                # If we have space, pad the last page or pad trailing bytes
                # Standard way: add spaces to last page if diff is even and >= 2
                if diff % 2 == 0 and diff > 0 and len(pages) > 0:
                    num_spaces = diff // 2
                    padded_last_page = pages[-1].strip() + (" " * num_spaces)
                    pages[-1] = padded_last_page
                    
                    new_s2_buf = bytearray()
                    new_s2_buf.append(len(pages))
                    for p in pages:
                        p_u16 = p.encode("utf-16le")
                        new_s2_buf.append(len(p_u16))
                        new_s2_buf.extend(p_u16)
                    
                # Write into dec at s2_start
                if len(new_s2_buf) == old_slot_len:
                    dec[s2_start:s2_start+old_slot_len] = new_s2_buf
                    updated_count += 1
                elif len(new_s2_buf) < old_slot_len:
                    dec[s2_start:s2_start+len(new_s2_buf)] = new_s2_buf
                    # zero out remainder
                    dec[s2_start+len(new_s2_buf):s2_start+old_slot_len] = b"\x00" * (old_slot_len - len(new_s2_buf))
                    updated_count += 1
            else:
                too_long_count += 1

        pos += rlen
        records_count += 1

    print(f"CPK Chapter 1 in-place patching results:")
    print(f"  Total records scanned: {records_count}")
    print(f"  Successfully patched in-place into CPK: {updated_count}")
    print(f"  Exceeded original slot length (kept original): {too_long_count}")

    # Encrypt back
    new_enc = encrypt_data(bytes(dec))
    assert len(new_enc) == file_size, f"Size changed! {len(new_enc)} != {file_size}"

    # Write back to parameter.cpk
    with open(cpk_path, "r+b") as f:
        f.seek(offset)
        f.write(new_enc)

    print(f"Successfully patched {file_size} bytes into parameter.cpk at offset {offset}!")
    print("parameter.cpk total size:", os.path.getsize(cpk_path))

if __name__ == "__main__":
    main()
