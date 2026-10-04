import openpyxl
import re

def main():
    wb = openpyxl.load_workbook(r"d:\Viet Hoa Game\I_am_Setsuna-parameter-Steam.xlsx", data_only=True)
    ws = wb["ScenarioMessageData_Chapter_1"]
    
    target_rids = [307, 1020, 1111, 1138, 1180, 1187, 1242, 1304, 1326, 1353]
    lines = []
    for r in range(3, 1406):
        id_val = ws.cell(r, 1).value
        m = re.search(r'"Id":(\d+)', str(id_val))
        if m and int(m.group(1)) in target_rids:
            rid = int(m.group(1))
            en = ws.cell(r, 2).value
            vn = ws.cell(r, 3).value
            lines.append(f"\n--- RID {rid} (Row {r}) ---")
            lines.append(f"EN ({len(str(en).encode('utf-16le'))} bytes): {repr(en)}")
            lines.append(f"VN ({len(str(vn).encode('utf-16le'))} bytes): {repr(vn)}")
            
    with open(r"d:\Viet Hoa Game\temp_scripts\long_vn_out.txt", "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print("Done! Check long_vn_out.txt")

if __name__ == "__main__":
    main()
